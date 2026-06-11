import json
import os
import uuid
from datetime import datetime
from collections import defaultdict
from typing import Any, List, Optional, Literal

from fastapi import APIRouter, Depends, UploadFile, File, Form, BackgroundTasks
from fastapi.responses import StreamingResponse
from langchain_core.messages import SystemMessage, HumanMessage, AIMessage, BaseMessage
from sqlalchemy.orm import Session, joinedload

import models, schemas
from database import get_db, SessionLocal
from auth import get_current_user
from agents import (
    task_extractor as task_agent,
    knowledge_manager as knowledge_agent,
    minutes_generator as minutes_agent,
    report_reviewer as report_agent,
)
from neo4j_client import get_meeting_graph_context, graph_context_to_str, run_cypher
from pydantic import BaseModel, Field
import logging

from .prompts import (
    make_llm,
    SUPERVISOR_DIRECT_SYSTEM, supervisor_direct_human,
)
from agent_logging import TokenUsageCollector, _token_collector_var, _create_log, _finalize

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/agent", tags=["agents"])


def _to_base_messages(messages: List[dict]) -> List[BaseMessage]:
    result = []
    for m in messages:
        role, content = m.get("role", ""), m.get("content", "") or ""
        if role == "user":
            result.append(HumanMessage(content=content))
        elif role in ("assistant", "agent"):
            result.append(AIMessage(content=content))
    return result



# ─── Knowledge Base 요청 스키마 ───────────────────────────────────────────────
class _StoreMinutesReq(BaseModel):
    meeting_id: int
    session_id: Optional[int] = None
    title: str
    content: str


class _StoreTaskReq(BaseModel):
    content: str
    department: Optional[str] = None
    due_date: Optional[str] = None
    meeting_id: Optional[int] = None


class _StoreReportReq(BaseModel):
    title: str
    content: str
    meeting_id: Optional[int] = None
    score: Optional[int] = None


class _ProposeRelationshipsReq(BaseModel):
    """POST /knowledge/propose-relationships 요청 바디."""
    meeting_id: int
    node_types: Optional[List[str]] = None  # None이면 Agenda·HumanJudgment·Minutes 전체


class _ConfirmRelationshipsReq(BaseModel):
    """POST /knowledge/confirm-relationships 요청 바디."""
    proposal_id: str
    approved: bool
    reject_reason: Optional[str] = None  # approved=False 일 때 반려 사유


# ─── Supervisor 라우팅 — LLM이 직접 에이전트를 선택 ───────────────────────────
class _RoutingDecision(BaseModel):
    """LLM 슈퍼바이저의 라우팅 결정 스트럭처드 아웃풋."""
    thinking: str = Field(
        description="어떤 에이전트가 적합한지 한국어로 1~2문장 근거 설명"
    )
    agent: Literal[
        "task_extractor",
        "minutes_generator",
        "report_reviewer",
        "knowledge_manager",
        "supervisor_direct",
    ] = Field(description="위임할 에이전트 이름")
    steps: List[str] = Field(
        default_factory=list,
        description=(
            "이 요청을 처리하기 위해 수행할 주요 작업을 한국어로 2~4단계 나열. "
            "각 항목은 짧은 한 문장(20자 이내), 번호·기호 없이."
        )
    )


_ROUTING_SYSTEM = """\
당신은 워크메이트 AI 슈퍼바이저입니다.
사용자의 요청을 분석하여 가장 적합한 에이전트를 선택하고, 처리 계획을 세우세요.

에이전트 선택 기준:
- task_extractor: 아젠다·과제·할 일·투두·Todo 새로 추출, 다음 회의 준비, 아카이브 파일 분석·추출
- minutes_generator: 회의록 작성·요약·편집, 회의 진행 보조, 실시간 통역·속기
- report_reviewer: 보고서·문서 검토·분석, 리뷰·피드백, 파일·자료 평가
- knowledge_manager: 과거 회의 내용 검색, 지식 베이스 저장·관리, HITL 검토·승인, 관계 그래프 조회
- supervisor_direct: 회의체 현황·브리핑, 과제 진행 상황 조회, 보고서 제출 현황 조회, 소속 회의체 목록, 구성원 안내, 인사·일반 질문
- off_topic: 회의체 운영과 전혀 관련 없는 질문 (날씨, 나이, 코딩, 개인 신상, 잡담, 일반 상식 등)

★ supervisor_direct 우선 케이스 (아래 패턴은 반드시 supervisor_direct):
  "브리핑", "현황", "상황 어때", "속해있어", "소속", "제출 현황", "진행 상황"

★ off_topic 케이스 예시 (반드시 off_topic):
  "너 몇살이야", "오늘 날씨 어때", "파이썬 코드 짜줘", "주식 어때", "점심 뭐 먹지", "농담 해줘"

thinking 필드에 선택 이유를 한국어 1~2문장으로 작성하세요.
steps 필드에 처리 계획을 한국어 2~4단계로 작성하세요. 각 단계는 20자 이내의 짧은 문장."""


async def classify_intent(message: str) -> tuple[str, str, List[str]]:
    """사용자 메시지를 분석해 (에이전트명, 근거, 처리단계) 튜플을 반환합니다."""
    try:
        routing_llm = make_llm(temperature=0.0, streaming=False).with_structured_output(_RoutingDecision)
        decision = await routing_llm.ainvoke([
            SystemMessage(content=_ROUTING_SYSTEM),
            HumanMessage(content=message[:500]),
        ])
        return decision.agent, decision.thinking, decision.steps or []
    except Exception as e:
        logger.warning(f"[Supervisor] 라우팅 LLM 실패, supervisor_direct 사용: {e}")
        return "supervisor_direct", "기본 처리 경로로 응답합니다.", []



# ─── Helpers ──────────────────────────────────────────────────────────────────
def _log_activity(meeting_id: int, agent: str, action: str, detail: str = ""):
    if not meeting_id:
        return
    db = SessionLocal()
    try:
        log = models.AgentLog(
            task_id=str(uuid.uuid4()),
            context_type=f"agent_{agent.lower()}",
            meeting_id=meeting_id,
            status="success",
            output_data={"agent": agent, "action": action, "detail": detail},
            ended_at=datetime.utcnow(),
        )
        db.add(log)
        db.commit()
    except Exception as e:
        db.rollback()
        print(f"[ActivityLog Error] {e}")
    finally:
        db.close()


def _get_meeting_context(db: Session, meeting_id: int) -> str:
    if not meeting_id:
        return ""
    meeting = db.query(models.Meeting).filter(models.Meeting.id == meeting_id).first()
    if not meeting:
        return ""
    lines = [f"회의체 이름: {meeting.title}"]
    if meeting.description:
        lines.append(f"회의 목적: {meeting.description}")
    members = db.query(models.MeetingMember).filter(
        models.MeetingMember.meeting_id == meeting_id
    ).all()
    if members:
        member_parts = []
        for m in members:
            user = db.query(models.User).filter(models.User.id == m.user_id).first()
            if user:
                role_label = "운영자" if m.role == "admin" else "참여자"
                dept = user.department or ""
                member_parts.append(f"{user.name}({dept}, {role_label})")
        if member_parts:
            lines.append(f"참여자: {', '.join(member_parts)}")
    return "\n".join(lines)


def _get_member_org_depts(db: Session, meeting_id: int) -> List[dict]:
    """Return unique (company, department) pairs from meeting members."""
    from sqlalchemy.orm import joinedload
    members = (
        db.query(models.MeetingMember)
        .options(joinedload(models.MeetingMember.user))
        .filter(models.MeetingMember.meeting_id == meeting_id)
        .all()
    )
    seen: set = set()
    result: List[dict] = []
    for m in members:
        if not m.user or not m.user.department:
            continue
        company = m.user.company or ""
        dept = m.user.department
        if (company, dept) not in seen:
            seen.add((company, dept))
            result.append({"company": company, "department": dept})
    return result


def _get_previous_minutes(db: Session, meeting_id: int) -> List[str]:
    sessions = db.query(models.MeetingSession).filter(
        models.MeetingSession.meeting_id == meeting_id,
        models.MeetingSession.status.in_(["ended", "ENDED"]),
    ).order_by(models.MeetingSession.ended_at.desc()).all()

    result = []
    for s in sessions:
        if not s.minutes:
            continue
        if s.minutes.content_summary:
            result.append(s.minutes.content_summary)
        elif s.minutes.file_path:
            # content_summary 없으면 파일에서 직접 텍스트 추출 (STT 원문 등)
            try:
                from r2_storage import is_r2_url as _is_r2, url_to_key as _r2_key, download_bytes as _r2_dl
                raw = _r2_dl(_r2_key(s.minutes.file_path)) if _is_r2(s.minutes.file_path) else None
                if raw:
                    text = _extract_text_from_file(raw, s.minutes.file_name or "minutes.pdf")
                    if text.strip():
                        result.append(text[:3000])
            except Exception:
                pass
    return result


def _format_schedule_table(table: list) -> str:
    """pdfplumber 테이블을 팀별 주차 일정 구조화 텍스트로 변환.

    pdfplumber가 셀을 세로로 분리해서 읽는 경우(팀명 행과 작업 행이 별도 행으로 추출)를
    처리하기 위해 row[0]이 None/빈값인 행을 직전 팀에 병합한다.
    """
    if not table or len(table) < 2:
        return ""

    def clean(cell) -> str:
        if cell is None:
            return ""
        return " ".join(str(cell).split())

    header = [clean(c) for c in table[0]]
    header_joined = " ".join(header)

    if not any(k in header_joined for k in ["주", "팀", "담당", "안건"]):
        return ""

    # 팀별로 col_idx → 작업명 매핑 (같은 팀의 여러 행을 병합)
    team_order: list = []
    team_tasks: dict = {}
    current_team: str | None = None

    for row in table[1:]:
        team_cell = clean(row[0]) if len(row) > 0 else ""
        if team_cell:
            current_team = team_cell
            if current_team not in team_tasks:
                team_order.append(current_team)
                team_tasks[current_team] = {}
        if current_team is None:
            continue
        for col_idx in range(1, len(row)):
            cell = clean(row[col_idx])
            if cell and col_idx not in team_tasks[current_team]:
                team_tasks[current_team][col_idx] = cell

    lines = ["[팀별 업무 일정표]"]
    for team in team_order:
        lines.append(f"[{team}]")
        tasks = team_tasks[team]
        if tasks:
            for col_idx in sorted(tasks):
                col_label = header[col_idx] if col_idx < len(header) and header[col_idx] else f"열{col_idx}"
                lines.append(f"  {col_label}: {tasks[col_idx]}")
        else:
            lines.append("  (배정된 작업 없음)")

    return "\n".join(lines) if len(lines) > 1 else ""


def _extract_text_from_file(raw: bytes, filename: str) -> str:
    import io

    if filename.endswith(".pdf"):
        # 1순위: pdfplumber — 표는 구조화 추출, 나머지는 텍스트 추출
        try:
            import pdfplumber
            parts = []
            with pdfplumber.open(io.BytesIO(raw)) as pdf:
                for page in pdf.pages:
                    # 표 구조화 추출
                    tables = page.extract_tables()
                    for t in (tables or []):
                        formatted = _format_schedule_table(t)
                        if formatted:
                            parts.append(formatted)
                    # 일반 텍스트 추출 (표 외 영역 포함)
                    page_text = page.extract_text() or ""
                    if page_text.strip():
                        parts.append(page_text.strip())
            result = "\n\n".join(parts).strip()
            if result:
                return result
        except ModuleNotFoundError:
            pass
        except Exception as e:
            logger.warning(f"[extract] pdfplumber 실패, pypdf로 대체: {e}")
        # 2순위: pypdf (기본 의존성)
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(raw))
            pages = [page.extract_text() or "" for page in reader.pages]
            return "\n".join(pages).strip()
        except Exception as e:
            return f"[PDF 추출 오류: {e}]"

    if filename.endswith(".docx"):
        try:
            import docx as _docx
            doc = _docx.Document(io.BytesIO(raw))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception as e:
            return f"[DOCX 추출 오류: {e}]"

    if filename.endswith((".xlsx", ".xls")):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    row_text = "\t".join(str(c) if c is not None else "" for c in row)
                    if row_text.strip():
                        lines.append(row_text)
            return "\n".join(lines)
        except Exception as e:
            return f"[XLSX 추출 오류: {e}]"

    for enc in ("utf-8", "euc-kr", "cp949"):
        try:
            return raw.decode(enc)
        except Exception:
            pass
    return ""


async def _stream_plan(system_content: str, human_content: str):
    """LLM이 작업 계획을 줄 단위로 스트리밍 생성합니다."""
    llm = make_llm(temperature=0.2, streaming=True)
    buf = ""
    async for chunk in llm.astream([
        SystemMessage(content=system_content),
        HumanMessage(content=human_content),
    ]):
        if chunk.content:
            buf += chunk.content
            while "\n" in buf:
                line, buf = buf.split("\n", 1)
                line = line.strip().lstrip("-•·▪▸◦*0123456789.").strip()
                if line:
                    yield line
    tail = buf.strip().lstrip("-•·▪▸◦*0123456789.").strip()
    if tail:
        yield tail


# ─── Minutes (아라) 에이전트 ──────────────────────────────────────────────────
@router.post("/minutes/sessions-chat")
async def minutes_sessions_chat(
    data: schemas.AgentChatRequest,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    sessions = (
        db.query(models.MeetingSession)
        .filter(models.MeetingSession.meeting_id == data.meeting_id)
        .order_by(models.MeetingSession.id.desc())
        .all()
    )
    sessions_info = [
        {
            "id": s.id,
            "title": s.title,
            "status": s.status,
            "summary": s.minutes.content_summary if s.minutes and s.minutes.content_summary else None,
        }
        for s in sessions
    ]
    agendas = db.query(models.Agenda).filter(
        models.Agenda.meeting_id == data.meeting_id,
    ).all()

    session_list_text = "\n".join([
        f"- {s['title']} ({s['status']})" + (": 요약 있음" if s['summary'] else ": 요약 없음")
        for s in sessions_info
    ])
    extra_context = f"[회의 세션 목록]\n{session_list_text}"
    session_summaries = [f"[{s['title']}] {s['summary']}" for s in sessions_info if s["summary"]]
    if session_summaries:
        extra_context += f"\n\n[세션별 회의록]\n" + "\n\n".join(session_summaries)

    async def stream():
        _collector = TokenUsageCollector()
        _tok_ctx_token = _token_collector_var.set(_collector)
        _log_id = _create_log(
            context_type="minutes_stream",
            meeting_id=data.meeting_id or None,
            session_id=None,
            user_id=current_user.id,
            input_data={"message": (data.message or "")[:300]},
        )
        _stream_error = None
        try:
            async for chunk in minutes_agent.chat_stream(
                message=data.message,
                chat_history=data.chat_history or [],
                previous_minutes=[extra_context],
                current_agendas=[{"content": a.title, "status": a.status} for a in agendas],
                meeting_context=_get_meeting_context(db, data.meeting_id),
            ):
                yield f"data: {chunk.replace(chr(10), chr(92)+chr(110))}\n\n"
            yield "data: [DONE]\n\n"
        except BaseException as _e:
            _stream_error = _e
            raise
        finally:
            _token_collector_var.reset(_tok_ctx_token)
            _finalize(_log_id, _collector, _stream_error, None)

    return StreamingResponse(stream(), media_type="text/event-stream")

@router.post("/minutes/generate-minutes")
async def minutes_generate_minutes(
    data: schemas.AgentChatRequest,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    transcript = data.message or ""
    meeting_context = _get_meeting_context(db, data.meeting_id) if data.meeting_id else ""
    agendas = db.query(models.Agenda).filter(models.Agenda.meeting_id == data.meeting_id).all() if data.meeting_id else []
    agenda_text = "\n".join([f"- {a.title} ({a.status})" for a in agendas]) or "없음"
    now = datetime.now().strftime("%Y년 %m월 %d일")
    meeting_obj = db.query(models.Meeting).filter(models.Meeting.id == data.meeting_id).first() if data.meeting_id else None
    minutes_title = f"{meeting_obj.title} 회의록 ({now})" if meeting_obj else f"회의록 ({now})"

    # session_info
    session_info = None
    if data.session_id:
        session = db.query(models.MeetingSession).filter(models.MeetingSession.id == data.session_id).first()
        if session:
            session_info = {
                "title": session.title,
                "started_at": session.started_at.strftime("%Y-%m-%d %H:%M") if session.started_at else None,
                "ended_at": session.ended_at.strftime("%Y-%m-%d %H:%M") if session.ended_at else None,
                "location": session.location,
            }

    # participants: session_id 있으면 세션 참석자, 없으면 회의체 멤버 전체
    participants = []
    if data.session_id:
        sm_rows = db.query(models.SessionMember).filter(models.SessionMember.session_id == data.session_id).all()
        user_ids = [sm.user_id for sm in sm_rows]
        role_map = {sm.user_id: sm.role for sm in sm_rows}
        users = db.query(models.User).filter(models.User.id.in_(user_ids)).all() if user_ids else []
        participants = [
            {"name": u.name, "dept": u.department or "", "role": role_map.get(u.id, "member")}
            for u in users
        ]
    elif data.meeting_id:
        mm_rows = (
            db.query(models.MeetingMember)
            .options(joinedload(models.MeetingMember.user))
            .filter(models.MeetingMember.meeting_id == data.meeting_id)
            .all()
        )
        participants = [
            {"name": mm.user.name, "dept": mm.user.department or "", "role": mm.role}
            for mm in mm_rows if mm.user
        ]

    async def stream():
        _collector = TokenUsageCollector()
        _tok_ctx_token = _token_collector_var.set(_collector)
        _log_id = _create_log(
            context_type="minutes_generate",
            meeting_id=data.meeting_id or None,
            session_id=data.session_id or None,
            user_id=current_user.id,
            input_data={"session_id": data.session_id},
        )
        _stream_error = None
        try:
            collected_parts = []
            async for chunk in minutes_agent.generate_minutes_stream(
                transcript, meeting_context, agenda_text, now,
                meeting_id=data.meeting_id,
                session_id=data.session_id,
                title=minutes_title,
            ):
                collected_parts.append(chunk)
                yield f"data: {chunk.replace(chr(10), chr(92)+chr(110))}\n\n"
            try:
                if data.session_id and collected_parts:
                    full_content = "".join(collected_parts)
                    _save_db = SessionLocal()
                    try:
                        existing = _save_db.query(models.Minutes).filter(
                            models.Minutes.session_id == data.session_id
                        ).first()
                        if existing:
                            existing.content_original = full_content
                            existing.content_summary = full_content[:500]
                            existing.recorder_id = current_user.id
                            existing.generated_at = datetime.utcnow()
                        else:
                            _save_db.add(models.Minutes(
                                session_id=data.session_id,
                                content_original=full_content,
                                content_summary=full_content[:500],
                                recorder_id=current_user.id,
                            ))
                        _save_db.commit()
                    except Exception as e:
                        logger.warning(f"[generate-minutes] PostgreSQL 저장 실패: {e}")
                        try: _save_db.rollback()
                        except Exception: pass
                    finally:
                        try: _save_db.close()
                        except Exception: pass
            except Exception as e:
                logger.warning(f"[generate-minutes] 저장 블록 예외: {e}")

            yield "data: [DONE]\n\n"
        except BaseException as _e:
            _stream_error = _e
            raise
        finally:
            _token_collector_var.reset(_tok_ctx_token)
            _finalize(_log_id, _collector, _stream_error, None)

    return StreamingResponse(stream(), media_type="text/event-stream")


# ─── Supervisor Chat ──────────────────────────────────────────────────────────
@router.post("/supervisor/chat")
async def supervisor_chat(
    data: schemas.AgentChatRequest,
    background_tasks: BackgroundTasks,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    msg = data.message or ""

    # ── LLM 라우팅 결정 ─────────────────────────────────────────────────────────
    _route, _route_thinking, _route_steps = await classify_intent(msg)

    background_tasks.add_task(
        _log_activity, data.meeting_id, f"워크메이트[{_route}]",
        "Supervisor 대화", f'"{msg[:80]}"'
    )

    user_person_id: str | None = None
    user_allowed_mg_ids: set[str] = set()
    is_admin = current_user.position in ("대표", "CEO", "임원")
    pg_meeting_ids: set[int] = {
        row.meeting_id
        for row in db.query(models.MeetingMember.meeting_id)
            .filter(models.MeetingMember.user_id == current_user.id)
            .all()
    }
    try:
        p_rows = await run_cypher(
            "MATCH (p:User) WHERE p.email = $email OR p.name = $name "
            "RETURN p.id AS pid LIMIT 1",
            {"email": current_user.email or "", "name": current_user.name or ""},
        )
        if p_rows:
            user_person_id = p_rows[0]["pid"]
            mg_access_rows = await run_cypher(
                "MATCH (p:User {id: $pid})-[:`간사`|`구성원`]->(mg) "
                "WHERE mg:Meetings OR mg:Meeting_session "
                "RETURN coalesce(mg.id, 'mg-sqlite-' + toString(mg.pg_id)) AS mg_id",
                {"pid": user_person_id},
            )
            user_allowed_mg_ids = {r["mg_id"] for r in mg_access_rows}
    except Exception:
        pass

    _thread_id = (
        data.thread_id
        if data.thread_id
        else (f"meeting_{data.meeting_id}" if data.meeting_id else f"global_{current_user.id}")
    )

    async def stream():
        _collector = TokenUsageCollector()
        _tok_ctx_token = _token_collector_var.set(_collector)
        _log_id = _create_log(
            context_type="supervisor",
            meeting_id=data.meeting_id or None,
            session_id=None,
            user_id=current_user.id,
            input_data={"message": msg[:300] if msg else None, "route": _route},
        )
        _stream_error = None

        print(f"DEBUG: stream() started, _route={_route!r}")
        # DB에서 최근 20개 대화 이력 조회 (시간 오름차순)
        _db_rows = (
            db.query(models.ChatMessage)
            .filter(
                models.ChatMessage.thread_id == _thread_id,
                models.ChatMessage.user_id == current_user.id,
            )
            .order_by(models.ChatMessage.created_at.desc())
            .limit(20)
            .all()
        )
        _chat_history_from_db: list[dict] = [
            {"role": m.role, "content": m.content or ""}
            for m in reversed(_db_rows)
        ]

        neo4j_ctx = {}
        neo4j_ctx_str = ""
        hl_candidates: list[str] = []
        try:
            print(f"DEBUG: outer try entered, meeting_id={data.meeting_id!r}")
            for _s in (_route_steps or [_route_thinking]):
                yield f"data: [PLANNING] {_s}\n\n"

            # ── off_topic 조기 종료: Neo4j·DB 조회 없이 안내 메시지 반환 ──────────
            if _route == 'off_topic':
                _off_msg = (
                    "저는 회의체 운영 관련 질문에 특화되어 있어요.\n"
                    "회의체 현황, 아젠다, 보고서, 회의록 관련 질문을 해주세요.\n\n"
                    "예를 들어:\n"
                    "- \"소속 회의체 현황 브리핑해줘\"\n"
                    "- \"아젠다 진행 상황 알려줘\"\n"
                    "- \"최근 보고서 제출 현황은?\""
                )
                yield f"data: {_off_msg.replace(chr(10), chr(92)+chr(110))}\n\n"
                yield "data: [DONE]\n\n"
                return

            print(f"DEBUG: after planning steps")
            if data.meeting_id:
                mid_neo4j = f"mg-{int(data.meeting_id)}"

                if not is_admin:
                    has_access = (
                        mid_neo4j in user_allowed_mg_ids if user_allowed_mg_ids
                        else int(data.meeting_id) in pg_meeting_ids
                    )
                    print(f"DEBUG: mid_neo4j={mid_neo4j!r}, user_allowed_mg_ids={user_allowed_mg_ids}, pg_meeting_ids={pg_meeting_ids}, has_access={has_access}")
                    if not has_access:
                        yield f"data: [PLANNING] 접근 권한 없음 — {current_user.name}님은 이 회의체에 대한 접근 권한이 없습니다\n\n"
                        yield "data: 이 회의체에 대한 접근 권한이 없습니다.\n\n"
                        yield "data: [DONE]\n\n"
                        return

                neo4j_ctx = await get_meeting_graph_context(data.meeting_id)

                if neo4j_ctx.get("meeting", {}).get("title"):
                    yield f"data: [PLANNING] [{neo4j_ctx['meeting']['title']}] 회의체 정보 확인\n\n"
                if neo4j_ctx.get("agendas"):
                    yield f"data: [PLANNING] 아젠다 {len(neo4j_ctx['agendas'])}건 분석\n\n"

                neo4j_ctx_str = graph_context_to_str(neo4j_ctx)

                for ag in neo4j_ctx.get("agendas", []):
                    if ag.get("title"): hl_candidates.append(ag["title"])
                for s in neo4j_ctx.get("recent_sessions", []):
                    if s.get("title"): hl_candidates.append(s["title"])
                if neo4j_ctx.get("meeting", {}).get("title"):
                    hl_candidates.append(neo4j_ctx["meeting"]["title"])
            else:
                try:
                    mg_detail_rows: list[dict] = []
                    if user_person_id:
                        mg_detail_rows = await run_cypher(
                            """MATCH (me:User {id: $pid})-[:`구성원`|`간사`]->(mg:Meetings)
                               WITH DISTINCT mg
                               OPTIONAL MATCH (sec:User)-[:`간사`]->(mg)
                               OPTIONAL MATCH (mem:User)-[:`구성원`]->(mg)
                               WITH mg,
                                    head(collect(DISTINCT
                                        sec.name + '||' + coalesce(sec.department, '')
                                    )) AS sec_info,
                                    [d IN collect(DISTINCT mem.department)
                                     WHERE d IS NOT NULL AND d <> ''] AS member_depts
                               OPTIONAL MATCH (s:Session)-[:`소속`]->(mg)
                               WITH mg, sec_info, member_depts,
                                    max(s.scheduled_at) AS latest_session_date
                               OPTIONAL MATCH (d:Report)-[:`첨부`]->(mg)
                               RETURN mg.id AS mg_id, mg.title AS title,
                                      coalesce(mg.meeting_type, '') AS meeting_type,
                                      sec_info, member_depts, latest_session_date,
                                      count(DISTINCT d) AS report_count
                               ORDER BY mg.title""",
                            {"pid": user_person_id},
                        )
                    elif is_admin:
                        mg_detail_rows = await run_cypher(
                            """MATCH (mg:Meetings)
                               OPTIONAL MATCH (sec:User)-[:`간사`]->(mg)
                               OPTIONAL MATCH (mem:User)-[:`구성원`]->(mg)
                               WITH mg,
                                    head(collect(DISTINCT
                                        sec.name + '||' + coalesce(sec.department, '')
                                    )) AS sec_info,
                                    [d IN collect(DISTINCT mem.department)
                                     WHERE d IS NOT NULL AND d <> ''] AS member_depts
                               OPTIONAL MATCH (s:Session)-[:`소속`]->(mg)
                               WITH mg, sec_info, member_depts,
                                    max(s.scheduled_at) AS latest_session_date
                               OPTIONAL MATCH (d:Report)-[:`첨부`]->(mg)
                               RETURN mg.id AS mg_id, mg.title AS title,
                                      coalesce(mg.meeting_type, '') AS meeting_type,
                                      sec_info, member_depts, latest_session_date,
                                      count(DISTINCT d) AS report_count
                               ORDER BY mg.title LIMIT 20"""
                        )

                    if mg_detail_rows:
                        # mg_id 기준 중복 제거
                        seen_mg_ids: set = set()
                        unique_rows: list[dict] = []
                        for _row in mg_detail_rows:
                            _mid = _row.get("mg_id", "")
                            if _mid and _mid in seen_mg_ids:
                                continue
                            if _mid:
                                seen_mg_ids.add(_mid)
                            unique_rows.append(_row)

                        yield f"data: [PLANNING] 소속 회의체 {len(unique_rows)}건 상세 조회\n\n"

                        ctx_lines = ["[소속 회의체 목록]"]
                        for _row in unique_rows:
                            _title = _row.get("title") or "?"
                            _mtype = _row.get("meeting_type") or ""
                            _sec_info = _row.get("sec_info") or ""
                            _depts: list = _row.get("member_depts") or []
                            _latest = _row.get("latest_session_date") or ""
                            _rcount = _row.get("report_count") or 0

                            _type_label = f" — {_mtype}" if _mtype else ""
                            ctx_lines.append(f"\n📋 {_title}{_type_label}")

                            if _sec_info:
                                _parts = _sec_info.split("||", 1)
                                _sec_name = _parts[0].strip()
                                _sec_dept = f"({_parts[1].strip()})" if len(_parts) > 1 and _parts[1].strip() else ""
                                ctx_lines.append(f"  - 간사: {_sec_name} {_sec_dept}".rstrip())
                            else:
                                ctx_lines.append("  - 간사: 미지정")

                            _unique_depts = list(dict.fromkeys(d for d in _depts if d))
                            ctx_lines.append(f"  - 참여부서: {', '.join(_unique_depts[:8])}" if _unique_depts else "  - 참여부서: 없음")

                            if _latest:
                                _date_str = str(_latest)[:10].replace("-", ".")
                                ctx_lines.append(f"  - 최근 회의: {_date_str}")
                            else:
                                ctx_lines.append("  - 최근 회의: 없음")

                            ctx_lines.append(f"  - 보고자료: {_rcount}건 제출")

                            if _title not in hl_candidates:
                                hl_candidates.append(_title)

                        neo4j_ctx_str = "\n".join(ctx_lines)

                    else:
                        # PostgreSQL fallback: Neo4j에 데이터 없는 경우
                        _pg_mids = list(pg_meeting_ids)[:20]
                        _pg_meetings = (
                            db.query(models.Meeting)
                            .filter(models.Meeting.id.in_(_pg_mids))
                            .order_by(models.Meeting.title)
                            .all()
                        ) if _pg_mids else []

                        if _pg_meetings:
                            yield f"data: [PLANNING] 소속 회의체 {len(_pg_meetings)}건 조회\n\n"
                            ctx_lines = ["[소속 회의체 목록]"]
                            for _mg in _pg_meetings:
                                _mems = (
                                    db.query(models.MeetingMember)
                                    .filter(models.MeetingMember.meeting_id == _mg.id)
                                    .all()
                                )
                                _user_ids = [m.user_id for m in _mems]
                                _users = {
                                    u.id: u for u in db.query(models.User)
                                    .filter(models.User.id.in_(_user_ids)).all()
                                }
                                _admin_mem = next((m for m in _mems if m.role == "admin"), None)
                                _sec = _users.get(_admin_mem.user_id) if _admin_mem else None
                                _depts = list(dict.fromkeys(
                                    _users[m.user_id].department
                                    for m in _mems
                                    if m.user_id in _users and _users[m.user_id].department
                                ))
                                _latest_s = (
                                    db.query(models.MeetingSession)
                                    .filter(
                                        models.MeetingSession.meeting_id == _mg.id,
                                        models.MeetingSession.status.in_(["ended", "ENDED"]),
                                    )
                                    .order_by(models.MeetingSession.ended_at.desc())
                                    .first()
                                )
                                _rcount = (
                                    db.query(models.Report)
                                    .filter(models.Report.meeting_id == _mg.id)
                                    .count()
                                )
                                _type_label = f" — {_mg.meeting_type}" if _mg.meeting_type else ""
                                ctx_lines.append(f"\n📋 {_mg.title}{_type_label}")
                                if _sec:
                                    _sec_dept = f"({_sec.department})" if _sec.department else ""
                                    ctx_lines.append(f"  - 간사: {_sec.name} {_sec_dept}".rstrip())
                                else:
                                    ctx_lines.append("  - 간사: 미지정")
                                ctx_lines.append(f"  - 참여부서: {', '.join(_depts[:8])}" if _depts else "  - 참여부서: 없음")
                                if _latest_s and _latest_s.ended_at:
                                    ctx_lines.append(f"  - 최근 회의: {_latest_s.ended_at.strftime('%Y.%m.%d')}")
                                else:
                                    ctx_lines.append("  - 최근 회의: 없음")
                                ctx_lines.append(f"  - 보고자료: {_rcount}건 제출")
                                if _mg.title not in hl_candidates:
                                    hl_candidates.append(_mg.title)
                            neo4j_ctx_str = "\n".join(ctx_lines)
                        else:
                            org_rows = await run_cypher(
                                "MATCH (org:Company) RETURN org.name AS name LIMIT 1"
                            )
                            if org_rows:
                                yield f"data: [PLANNING] 조직: {org_rows[0].get('name', '?')} 확인\n\n"
                except Exception:
                    pass

        except Exception as _outer_e:
            print(f"DEBUG: outer except caught: {type(_outer_e).__name__}: {_outer_e}")
            yield "data: [PLANNING] 지식 그래프 조회 중 오류 발생\n\n"

        _user_scope_header = (
            f"[현재 사용자] {current_user.name}"
            + (f" / {current_user.position}" if current_user.position else "")
            + (f" / {current_user.department}" if current_user.department else "")
            + "\n[데이터 접근 범위] 본인이 소속된 회의체의 정보만 제공합니다. "
            "다른 사용자 또는 비소속 회의체의 민감 정보는 노출하지 마세요.\n"
        )

        def _enrich(base_ctx: str) -> str:
            parts = [_user_scope_header, base_ctx]
            if neo4j_ctx_str and neo4j_ctx_str != "(Neo4j 데이터 없음)":
                parts.append(f"[Neo4j 그래프 컨텍스트]\n{neo4j_ctx_str}")
            return "\n\n".join(parts)

        # ── 라우팅 공통: assistant 응답 수집 (finally에서 저장) ─────────────
        _assistant_chunks: list[str] = []

        try:
            # ── B 유형: 현황 조회 / 지식 베이스 / 인사 / 일반 질문 ──────────
            print(f"DEBUG: routing block entered, _route={_route!r}")

            if _route in ('supervisor_direct', 'knowledge_manager'):
                yield "data: [PLANNING] Knowledge Base에서 관련 자료 검색 중...\n\n"

                _kb_results: list[dict] = []
                _node_type_map = [("Minutes", "회의록", 5)]
                if _route == "knowledge_manager":
                    _node_type_map += [("Agenda", "안건", 3)]
                try:
                    for node_type, type_label, k in _node_type_map:
                        for r in await knowledge_agent.search_knowledge(msg, node_type=node_type, k=k):
                            if r.get("title") or r.get("content"):
                                _kb_results.append({
                                    "type": type_label,
                                    "title": r.get("title") or r.get("content", "")[:40],
                                    "content": r.get("content", "")[:300],
                                })
                except Exception:
                    pass

                if _kb_results:
                    yield f"data: [PLANNING] 관련 자료 {len(_kb_results)}건 확인\n\n"

                _ctx_parts: list[str] = [_user_scope_header]
                if neo4j_ctx_str and neo4j_ctx_str != "(Neo4j 데이터 없음)":
                    _ctx_parts.append(f"[회의체 현황]\n{neo4j_ctx_str}")

                # Neo4j에 멤버 데이터가 없을 경우 PostgreSQL로 보완
                if data.meeting_id and not neo4j_ctx.get("members"):
                    _pg_ctx = _get_meeting_context(db, data.meeting_id)
                    if _pg_ctx:
                        _ctx_parts.append(f"[회의체 기본 정보]\n{_pg_ctx}")

                # 최근 AgentLog 활동 추가 (있을 경우)
                if data.meeting_id:
                    try:
                        _recent_logs = (
                            db.query(models.AgentLog)
                            .filter(models.AgentLog.meeting_id == data.meeting_id)
                            .order_by(models.AgentLog.ended_at.desc())
                            .limit(5)
                            .all()
                        )
                        if _recent_logs:
                            _log_lines = []
                            for _log in _recent_logs:
                                _detail = (_log.output_data or {}).get("action", "") or (_log.output_data or {}).get("detail", "")
                                _log_lines.append(f"  - {_log.context_type}: {_detail[:60]}" if _detail else f"  - {_log.context_type}")
                            _ctx_parts.append("[최근 활동 로그]\n" + "\n".join(_log_lines))
                    except Exception:
                        pass

                # ── 과제 진행 상황 — 회의체별 집계 (PostgreSQL) ──────────────
                try:
                    _agenda_meeting_ids = (
                        [data.meeting_id] if data.meeting_id
                        else [mid for mid in pg_meeting_ids][:10]
                    )
                    if _agenda_meeting_ids:
                        _agendas = (
                            db.query(models.Agenda)
                            .filter(
                                models.Agenda.meeting_id.in_(_agenda_meeting_ids),
                            )
                            .order_by(models.Agenda.meeting_id, models.Agenda.due_date)
                            .all()
                        )
                        if _agendas:
                            from collections import Counter as _Counter, defaultdict as _dd_a
                            # 회의체별 status 카운트
                            _by_mg_agenda: dict = _dd_a(list)
                            for _a in _agendas:
                                _by_mg_agenda[_a.meeting_id].append(_a)

                            # 회의체 title 캐시
                            _mg_title_cache: dict[int, str] = {}
                            for _mid in _agenda_meeting_ids:
                                _m = db.query(models.Meeting).filter(models.Meeting.id == _mid).first()
                                if _m:
                                    _mg_title_cache[_mid] = _m.title

                            _a_lines = [f"[아젠다 현황] 총 {len(_agendas)}건 (회의체별 집계)"]
                            for _mid, _alist in _by_mg_agenda.items():
                                _mg_name = _mg_title_cache.get(_mid, f"회의체 {_mid}")
                                _cnt = _Counter(a.status for a in _alist)
                                _a_lines.append(f"\n  [{_mg_name}]")
                                for _s, _label in [
                                    ("ongoing", "진행 중"),
                                    ("done", "완료"),
                                    ("pending", "대기"),
                                    ("submitted", "제출완료"),
                                    ("draft", "초안"),
                                ]:
                                    if _cnt.get(_s):
                                        _a_lines.append(f"    - {_label}: {_cnt[_s]}건")
                            _ctx_parts.append("\n".join(_a_lines))
                except Exception:
                    pass

                # ── 보고서 제출 현황 (PostgreSQL) ─────────────────────────────
                try:
                    _report_meeting_ids = (
                        [data.meeting_id] if data.meeting_id
                        else [mid for mid in pg_meeting_ids][:10]
                    )
                    if _report_meeting_ids:
                        _reports = (
                            db.query(models.Report)
                            .filter(models.Report.meeting_id.in_(_report_meeting_ids))
                            .order_by(models.Report.created_at.desc())
                            .limit(30)
                            .all()
                        )
                        if _reports:
                            _mg_titles: dict[int, str] = {}
                            for _mid in _report_meeting_ids:
                                _m = db.query(models.Meeting).filter(models.Meeting.id == _mid).first()
                                if _m:
                                    _mg_titles[_mid] = _m.title
                            _r_status_label = {
                                "pending": "검토중", "approved": "승인",
                                "rejected": "반려", "draft": "초안",
                            }
                            from collections import defaultdict as _dd
                            _by_mg: dict = _dd(list)
                            for _r in _reports:
                                _by_mg[_r.meeting_id].append(_r)
                            _r_lines = [f"[보고서 제출 현황] 총 {len(_reports)}건"]
                            for _mid, _rlist in _by_mg.items():
                                _mg_name = _mg_titles.get(_mid, f"회의체 {_mid}")
                                _r_lines.append(f"\n  [{_mg_name}] {len(_rlist)}건")
                                for _r in _rlist[:5]:
                                    _rs = _r_status_label.get(_r.human_status or "", _r.human_status or "")
                                    _rdate = _r.created_at.strftime("%Y.%m.%d") if _r.created_at else ""
                                    _dept = _r.submitter_department or "미상"
                                    _r_lines.append(
                                        f"    - {_r.file_name or '(파일없음)'} "
                                        f"[{_dept}] [{_rs}] ({_rdate})"
                                    )
                            _ctx_parts.append("\n".join(_r_lines))
                except Exception:
                    pass

                if _kb_results:
                    _ctx_parts.append(
                        "[Knowledge Base 관련 자료]\n" + "\n".join(
                            f"- [{r['type']}] {r['title']}: {r['content'][:200]}" for r in _kb_results
                        )
                    )

                _sup_llm = make_llm(temperature=0.2, streaming=True)
                _history_msgs = _to_base_messages(_chat_history_from_db)
                async for chunk in _sup_llm.astream(
                    [SystemMessage(content=SUPERVISOR_DIRECT_SYSTEM)]
                    + _history_msgs
                    + [HumanMessage(content=supervisor_direct_human(msg, "\n\n".join(_ctx_parts)))]
                ):
                    if chunk.content:
                        _assistant_chunks.append(chunk.content)
                        yield f"data: {chunk.content.replace(chr(10), chr(92)+chr(110))}\n\n"

                if hl_candidates and _assistant_chunks:
                    _sup_full = "".join(_assistant_chunks)
                    matched = [c for c in hl_candidates if c and c in _sup_full]
                    if matched:
                        yield f"data: [HIGHLIGHT] {json.dumps(matched, ensure_ascii=False)}\n\n"

                yield "data: [DONE]\n\n"
                return

            # ── A 유형: 서브에이전트 라우팅 ──────────────────────────────────
            if _route == 'task_extractor':
                _org_dept_pairs = _get_member_org_depts(db, data.meeting_id)
                _org_dept_list = (
                    "\n".join(
                        f"- {p['company']} / {p['department']}" if p.get("company") else f"- {p['department']}"
                        for p in _org_dept_pairs
                    ) if _org_dept_pairs else "정보 없음"
                )
                gen = task_agent.chat_stream(
                    message=msg, chat_history=_chat_history_from_db,
                    previous_minutes=_get_previous_minutes(db, data.meeting_id),
                    knowledge=[],
                    org_dept_list=_org_dept_list,
                    meeting_context=_enrich(_get_meeting_context(db, data.meeting_id)),
                )
            elif _route == 'minutes_generator':
                agendas = db.query(models.Agenda).filter(
                    models.Agenda.meeting_id == data.meeting_id,
                    models.Agenda.status.in_(["ON_HOLD", "IN_PROGRESS"]),
                ).all()
                gen = minutes_agent.chat_stream(
                    message=msg, chat_history=_chat_history_from_db,
                    previous_minutes=_get_previous_minutes(db, data.meeting_id),
                    current_agendas=[{'content': a.title, 'status': a.status} for a in agendas],
                    meeting_context=_enrich(_get_meeting_context(db, data.meeting_id)),
                )
            else:  # report_reviewer
                gen = report_agent.chat_stream(
                    message=msg, chat_history=_chat_history_from_db,
                    knowledge=[],
                    meeting_context=_enrich(_get_meeting_context(db, data.meeting_id)),
                )

            async for chunk in gen:
                _assistant_chunks.append(chunk)
                yield f"data: {chunk.replace(chr(10), chr(92)+chr(110))}\n\n"

            if hl_candidates and _assistant_chunks:
                full_text = "".join(_assistant_chunks)
                matched = [c for c in hl_candidates if c and c in full_text]
                if matched:
                    yield f"data: [HIGHLIGHT] {json.dumps(matched, ensure_ascii=False)}\n\n"

            yield "data: [DONE]\n\n"

        except BaseException as _e:
            _stream_error = _e
            raise
        finally:
            # GeneratorExit·예외·정상 종료 모두에서 assistant 메시지 저장
            _token_collector_var.reset(_tok_ctx_token)
            _finalize(_log_id, _collector, _stream_error, None)
            _assistant_text = "".join(_assistant_chunks)
            if _assistant_text:
                _save_db = SessionLocal()
                try:
                    _save_db.add(models.ChatMessage(
                        thread_id=_thread_id,
                        user_id=current_user.id,
                        role="assistant",
                        content=_assistant_text,
                        context_type="supervisor",
                        meeting_id=data.meeting_id or None,
                    ))
                    _save_db.commit()
                except Exception as _e:
                    logger.warning(f"[supervisor_chat] AI 응답 저장 실패: {_e}")
                    _save_db.rollback()
                finally:
                    _save_db.close()

    # 사용자 메시지 저장
    if msg:
        try:
            db.add(models.ChatMessage(
                thread_id=_thread_id,
                user_id=current_user.id,
                role="user",
                content=msg,
                context_type="supervisor",
                meeting_id=data.meeting_id or None,
            ))
            db.commit()
        except Exception as _e:
            logger.warning(f"[supervisor_chat] 사용자 메시지 저장 실패: {_e}")
            db.rollback()

    return StreamingResponse(stream(), media_type="text/event-stream")


# ─── Supervisor Chat 히스토리 조회 ───────────────────────────────────────────
@router.get("/supervisor/chat/history")
async def supervisor_chat_history(
    meeting_id: int,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    is_admin = current_user.position in ("대표", "CEO", "임원")
    if not is_admin:
        member = db.query(models.MeetingMember).filter(
            models.MeetingMember.meeting_id == meeting_id,
            models.MeetingMember.user_id == current_user.id,
        ).first()
        if not member:
            from fastapi import HTTPException
            raise HTTPException(status_code=403, detail="이 회의체에 대한 접근 권한이 없습니다.")

    messages = (
        db.query(models.ChatMessage)
        .filter(
            models.ChatMessage.meeting_id == meeting_id,
            models.ChatMessage.context_type == "supervisor",
        )
        .order_by(models.ChatMessage.created_at.asc())
        .all()
    )
    return [
        {
            "id": m.id,
            "role": m.role,
            "content": m.content,
            "user_id": m.user_id,
            "created_at": m.created_at.isoformat() if m.created_at else None,
        }
        for m in messages
    ]


# ─── Supervisor 그래프 분석 ───────────────────────────────────────────────────
_AGENDA_LINK_THRESHOLD = 0.80
_DOC_LINK_THRESHOLD    = 0.78


async def _count_nodes(label: str) -> int:
    try:
        r = await run_cypher(f"MATCH (n:{label}) RETURN count(n) AS c")
        return r[0]["c"] if r else 0
    except Exception:
        return 0


async def _analyze_graph() -> dict:
    semantic   = {"agenda_links": [], "doc_links": []}
    structural = {"session_chains": [], "lifecycle_gaps": [], "stale_links": [], "ownerless_agendas": [],
                  "orphan_documents": [], "minuteless_sessions": [], "isolated_persons": []}
    membership = {"issues": []}

    counts = {
        "agendas":   await _count_nodes("Agenda"),
        "documents": (await _count_nodes("Report")) + (await _count_nodes("Minutes")),
        "meetings":  await _count_nodes("Meetings"),
        "persons":   await _count_nodes("User"),
        "sessions":  await _count_nodes("Session"),
    }

    # ⓪ 세션 시간순 체인 점검
    try:
        ch_rows = await run_cypher(
            "MATCH (s:Session)-[:`소속`|`개최`]->(mg:Meetings) "
            "WITH mg, s ORDER BY CASE WHEN coalesce(s.scheduled_at,'')='' THEN 1 ELSE 0 END, "
            "     s.scheduled_at, s.id "
            "WITH mg, collect({id:s.id, title:coalesce(s.title,'')}) AS sess "
            "WHERE size(sess) >= 2 "
            "RETURN mg.id AS mg_id, coalesce(mg.title,'') AS mg_title, sess AS ordered"
        )
        existing_followups: set = set()
        try:
            ex_rows = await run_cypher(
                "MATCH (a:Session)-[:`후속`]->(b:Session) RETURN a.id AS a, b.id AS b"
            )
            existing_followups = {(r.get("a"), r.get("b")) for r in ex_rows}
        except Exception:
            pass
        for r in ch_rows:
            ordered = r.get("ordered", []) or []
            missing = [
                {"a_id": ordered[i].get("id"), "a_title": ordered[i].get("title", ""),
                 "b_id": ordered[i+1].get("id"), "b_title": ordered[i+1].get("title", "")}
                for i in range(len(ordered) - 1)
                if (ordered[i].get("id"), ordered[i+1].get("id")) not in existing_followups
            ]
            if missing:
                structural["session_chains"].append(
                    {"mg": r.get("mg_title", ""), "count": len(ordered), "missing": missing}
                )
    except Exception as e:
        logger.warning(f"[Supervisor] 세션 체인 분석 실패(무시): {e}")

    # ⓪-b 회의 생명주기 공백
    try:
        lg_rows = await run_cypher(
            "MATCH (mn:Minutes)-[:`생성`]->(s:Session)-[:`소속`|`개최`]->(mg:Meetings) "
            "WHERE coalesce(mn.content_summary,'') <> '' AND NOT (mn)-[:`도출`]->(:Agenda) "
            "OPTIONAL MATCH (ag:Agenda)-[:`관할`]->(mg) "
            "WITH mn, s, mg, collect(DISTINCT {id: coalesce(ag.id, toString(ag.pg_id)), title: ag.title})[..25] AS ags "
            "RETURN mn.pg_id AS minutes_pg_id, s.id AS session_id, coalesce(s.title,'') AS session_title, "
            "       coalesce(mg.title,'') AS mg, left(mn.content_summary, 1800) AS content, ags AS agendas "
            "LIMIT 10"
        )
        for r in lg_rows:
            structural["lifecycle_gaps"].append({
                "minutes_pg_id": r.get("minutes_pg_id"),
                "session_id": r.get("session_id"),
                "session_title": r.get("session_title", ""),
                "mg": r.get("mg", ""),
                "content": r.get("content", ""),
                "agendas": [a for a in (r.get("agendas") or []) if a.get("title")],
            })
    except Exception as e:
        logger.warning(f"[Supervisor] 생명주기 공백 분석 실패(무시): {e}")

    # ① 잠재 연결 — 회의 간 의미 유사 안건쌍
    try:
        rows = await run_cypher(
            "MATCH (a:Agenda)-[:`관할`]->(mga:Meetings) "
            "WHERE a.embedding IS NOT NULL "
            "CALL db.index.vector.queryNodes('agendaEmbedding', 5, a.embedding) "
            "YIELD node, score "
            "WITH a, mga, node, score "
            "WHERE a.id < node.id AND score >= $th AND node.embedding IS NOT NULL "
            "MATCH (node)-[:`관할`]->(mgb:Meetings) "
            "WHERE mgb.id <> mga.id AND NOT (a)-[:`관련`]-(node) "
            "RETURN a.id AS a_id, a.title AS a_title, mga.title AS a_mg, "
            "       node.id AS b_id, node.title AS b_title, mgb.title AS b_mg, score "
            "ORDER BY score DESC LIMIT 30",
            {"th": _AGENDA_LINK_THRESHOLD},
        )
        for r in rows:
            semantic["agenda_links"].append({
                "a_id": r.get("a_id"), "a_title": r.get("a_title", "?"), "a_mg": r.get("a_mg", ""),
                "b_id": r.get("b_id"), "b_title": r.get("b_title", "?"), "b_mg": r.get("b_mg", ""),
                "score": float(r.get("score") or 0.0),
            })
    except Exception as e:
        logger.warning(f"[Supervisor] 안건 유사도 분석 실패(무시): {e}")

    # ① 잠재 연결 — 문서 ↔ 안건 적합도
    try:
        rows = await run_cypher(
            "MATCH (d) WHERE (d:Report OR d:Minutes) AND d.embedding IS NOT NULL "
            "CALL db.index.vector.queryNodes('agendaEmbedding', 3, d.embedding) "
            "YIELD node, score "
            "WITH d, node, score "
            "WHERE score >= $th AND NOT (d)-[:`첨부`]->(node) AND NOT (d)-[:`참조`]->(node) "
            "RETURN d.id AS doc_id, coalesce(d.title, d.file_name) AS doc_title, "
            "       node.id AS ag_id, node.title AS ag_title, score "
            "ORDER BY score DESC LIMIT 30",
            {"th": _DOC_LINK_THRESHOLD},
        )
        for r in rows:
            semantic["doc_links"].append({
                "doc_id": r.get("doc_id"), "doc_title": r.get("doc_title", "?"),
                "ag_id": r.get("ag_id"), "ag_title": r.get("ag_title", "?"),
                "score": float(r.get("score") or 0.0),
            })
    except Exception as e:
        logger.warning(f"[Supervisor] 문서-안건 적합도 분석 실패(무시): {e}")

    # ② 구조 공백 — 담당자 없는 안건
    try:
        rows = await run_cypher(
            "MATCH (ag:Agenda) WHERE NOT (:User)-[:`담당`]->(ag) "
            "OPTIONAL MATCH (ag)-[:`관할`]->(mg:Meetings) "
            "RETURN ag.id AS id, ag.title AS title, mg.title AS mg LIMIT 50"
        )
        structural["ownerless_agendas"] = [
            {"id": r.get("id"), "title": r.get("title", "?"), "mg": r.get("mg", "")} for r in rows
        ]
    except Exception:
        pass

    # ② 구조 공백 — 고아 문서
    try:
        rows = await run_cypher(
            "MATCH (d:Report) WHERE NOT (d)-[:`첨부`]->() "
            "RETURN d.id AS id, coalesce(d.title, d.file_name) AS title, "
            "       (d.embedding IS NOT NULL) AS emb LIMIT 50"
        )
        structural["orphan_documents"] = [
            {"id": r.get("id"), "title": r.get("title", "?"), "emb": bool(r.get("emb"))} for r in rows
        ]
    except Exception:
        pass

    # ② 구조 공백 — 회의록 없는 세션
    try:
        rows = await run_cypher(
            "MATCH (s:Session) WHERE NOT (:Minutes)-[:`생성`]->(s) "
            "OPTIONAL MATCH (s)-[:`소속`]->(mg:Meetings) "
            "RETURN s.id AS id, mg.title AS mg LIMIT 50"
        )
        structural["minuteless_sessions"] = [
            {"id": r.get("id"), "mg": r.get("mg", "")} for r in rows
        ]
    except Exception:
        pass

    # ② 구조 공백 — 고립 인물
    try:
        rows = await run_cypher(
            "MATCH (p:User) WHERE NOT (p)--() RETURN p.id AS id, p.name AS name LIMIT 50"
        )
        structural["isolated_persons"] = [
            {"id": r.get("id"), "name": r.get("name", "?")} for r in rows
        ]
    except Exception:
        pass

    # ③ 소속 무결성
    for cypher, issue_type, extra_key in [
        (
            "MATCH (p:User)-[:`소속부서`]->(d:Department) "
            "RETURN p.id AS pid, p.name AS name, d.name AS dept",
            "legacy", None,
        ),
        (
            "MATCH (p:User) WHERE coalesce(p.department, '') <> '' "
            "  AND NOT (p)-[:`소속`]->(:Department) "
            "RETURN p.id AS pid, p.name AS name, p.department AS dept",
            "missing", None,
        ),
        (
            "MATCH (p:User)-[:`소속`]->(d:Department) "
            "WHERE coalesce(p.department, '') <> '' AND d.name <> p.department "
            "RETURN p.id AS pid, p.name AS name, p.department AS dept, d.name AS wrong",
            "mismatch", "wrong",
        ),
    ]:
        try:
            for r in await run_cypher(cypher):
                entry = {"type": issue_type, "pid": r.get("pid"), "person": r.get("name", "?"),
                         "dept": r.get("dept", ""), "current": r.get(extra_key) if extra_key else None}
                membership["issues"].append(entry)
        except Exception:
            pass

    # ① 불필요 연결 탐지
    _PRUNE_THRESHOLD = 0.70

    # stale carry-forward: 완료된 안건에 달린 세션 도출(이월) 관계
    try:
        rows = await run_cypher(
            "MATCH (s:Session)-[r:`도출`]->(ag:Agenda) "
            "WHERE r.kind = 'carry_forward' "
            "  AND ag.status IN ['DONE', 'COMPLETED', 'CLOSED', 'RESOLVED'] "
            "RETURN s.id AS session_id, ag.id AS agenda_id, ag.title AS agenda_title"
        )
        for r in rows:
            structural["stale_links"].append({
                "kind": "stale_carry",
                "from_id": r.get("session_id"),
                "to_id": r.get("agenda_id"),
                "label": f"이월 → [{r.get('agenda_title','?')}] (완료된 안건)",
                "rel": "도출",
            })
    except Exception:
        pass

    # stale lifecycle: 완료된 안건에 달린 Minutes 도출 관계
    try:
        rows = await run_cypher(
            "MATCH (mn:Minutes)-[r:`도출`]->(ag:Agenda) "
            "WHERE r.kind = 'minutes_agenda' "
            "  AND ag.status IN ['DONE', 'COMPLETED', 'CLOSED', 'RESOLVED'] "
            "RETURN mn.pg_id AS mn_id, ag.id AS agenda_id, ag.title AS agenda_title"
        )
        for r in rows:
            structural["stale_links"].append({
                "kind": "stale_lifecycle",
                "from_id": str(r.get("mn_id", "")),
                "to_id": r.get("agenda_id"),
                "label": f"회의록도출 → [{r.get('agenda_title','?')}] (완료된 안건)",
                "rel": "도출",
            })
    except Exception:
        pass

    # weak 관련 (낙은 유사도 자동 생성 안건시안 링크)
    try:
        rows = await run_cypher(
            "MATCH (a:Agenda)-[r:`관련`]-(b:Agenda) "
            "WHERE r.discovered_by = 'knowledge_agent' AND r.score IS NOT NULL AND r.score < $th "
            "RETURN a.id AS from_id, b.id AS to_id, a.title AS a_title, b.title AS b_title, r.score AS score "
            "LIMIT 30",
            {"th": _PRUNE_THRESHOLD},
        )
        for r in rows:
            structural["stale_links"].append({
                "kind": "weak_related",
                "from_id": r.get("from_id"),
                "to_id": r.get("to_id"),
                "label": f"[{r.get('a_title','?')}]↔[{r.get('b_title','?')}] ({float(r.get('score') or 0)*100:.0f}%)",
                "rel": "관련",
                "score": float(r.get("score") or 0),
            })
    except Exception:
        pass

    # weak 참조 (낙은 유사도 자동 문서-안건 연결)
    try:
        rows = await run_cypher(
            "MATCH (d)-[r:`참조`]->(ag:Agenda) "
            "WHERE (d:Report OR d:Minutes) AND r.discovered_by = 'knowledge_agent' AND r.score IS NOT NULL AND r.score < $th "
            "RETURN d.id AS from_id, ag.id AS to_id, "
            "       coalesce(d.title, d.file_name, '?') AS doc_title, ag.title AS ag_title, r.score AS score "
            "LIMIT 30",
            {"th": _PRUNE_THRESHOLD},
        )
        for r in rows:
            structural["stale_links"].append({
                "kind": "weak_ref",
                "from_id": r.get("from_id"),
                "to_id": r.get("to_id"),
                "label": f"문서[{r.get('doc_title','?')}]→안건[{r.get('ag_title','?')}] ({float(r.get('score') or 0)*100:.0f}%)",
                "rel": "참조",
                "score": float(r.get("score") or 0),
            })
    except Exception:
        pass

    # weak 첨부 (낙은 유사도 자동 첨부 연결)
    try:
        rows = await run_cypher(
            "MATCH (d)-[r:`첨부`]->(ag:Agenda) "
            "WHERE (d:Report OR d:Minutes) AND r.auto_linked = true AND r.score IS NOT NULL AND r.score < $th "
            "RETURN d.id AS from_id, ag.id AS to_id, "
            "       coalesce(d.title, d.file_name, '?') AS doc_title, ag.title AS ag_title, r.score AS score "
            "LIMIT 30",
            {"th": _PRUNE_THRESHOLD},
        )
        for r in rows:
            structural["stale_links"].append({
                "kind": "weak_attach",
                "from_id": r.get("from_id"),
                "to_id": r.get("to_id"),
                "label": f"문서[{r.get('doc_title','?')}]→안건[{r.get('ag_title','?')}] 첨부 ({float(r.get('score') or 0)*100:.0f}%)",
                "rel": "첨부",
                "score": float(r.get("score") or 0),
            })
    except Exception:
        pass

    return {"semantic": semantic, "structural": structural,
            "membership": membership, "counts": counts}


async def _normalize_rel_directions() -> dict:
    """관계 방향·명칭을 정규 흐름(REL_MATRIX)에 맞게 일괄 정규화합니다.
    Neo4j는 엣지 방향을 in-place 변경할 수 없으므로 DELETE → MERGE 패턴을 사용합니다.
    """
    # 원칙: 포함/소속은 작은 단위 → 큰 단위 방향 (child → parent)
    # 라이프사이클(agenda→session→minutes→human_judgment)은 흐름 방향 유지
    steps = [
        # (설명, 카운트 쿼리, 수정 쿼리)

        # ── 포함 관계: 잘못된 big→small 복구 ──────────────────────────
        ("회의체→회사 방향 복구 (포함, 작은→큰)",
         "MATCH (co:Company)-[r:`포함`]->(mg:Meetings) RETURN count(r) AS cnt",
         "MATCH (co:Company)-[r:`포함`]->(mg:Meetings) MERGE (mg)-[:`포함`]->(co) DELETE r"),

        ("부서→회사 방향 복구 (소속, 작은→큰)",
         "MATCH (co:Company)-[r:`소속`|`포함`]->(d:Department) RETURN count(r) AS cnt",
         "MATCH (co:Company)-[r:`소속`|`포함`]->(d:Department) MERGE (d)-[:`소속`]->(co) DELETE r"),

        ("사람→부서 방향 복구 (소속, 작은→큰)",
         "MATCH (d:Department)-[r:`소속`]->(u:User) RETURN count(r) AS cnt",
         "MATCH (d:Department)-[r:`소속`]->(u:User) MERGE (u)-[:`소속`]->(d) DELETE r"),

        ("아젠다→회의체 방향 복구 (관할, 작은→큰)",
         "MATCH (mg:Meetings)-[r:`관할`]->(ag:Agenda) RETURN count(r) AS cnt",
         "MATCH (mg:Meetings)-[r:`관할`]->(ag:Agenda) MERGE (ag)-[:`관할`]->(mg) DELETE r"),

        # ── 라이프사이클 방향 정규화 ───────────────────────────────────
        ("회의록→의사결정 방향 정규화 (판단)",
         "MATCH (hj:HumanJudgment)-[r:`판단`]->(n) RETURN count(r) AS cnt",
         "MATCH (hj:HumanJudgment)-[r:`판단`]->(n) MERGE (n)-[:`판단`]->(hj) DELETE r"),

        ("'다룸멌' 명칭·방향 정규화 → 아젠다→회의 (다룸)",
         "MATCH (s:Session)-[r:`다룸멌`]->(ag:Agenda) RETURN count(r) AS cnt",
         "MATCH (s:Session)-[r:`다룸멌`]->(ag:Agenda) MERGE (ag)-[:`다룸`]->(s) DELETE r"),

        ("회의→아젠다 '다룸' 방향 정규화",
         "MATCH (s:Session)-[r:`다룸`]->(ag:Agenda) RETURN count(r) AS cnt",
         "MATCH (s:Session)-[r:`다룸`]->(ag:Agenda) MERGE (ag)-[:`다룸`]->(s) DELETE r"),

        # ── 누락 연결 보완 ─────────────────────────────────────────────
        ("부서-회사 누락 연결 보완 (소속, 작은→큰)",
         "MATCH (d:Department) WHERE NOT (d)-[:`소속`]->(:Company) RETURN count(d) AS cnt",
         "MATCH (d:Department), (co:Company) WHERE NOT (d)-[:`소속`]->(co) MERGE (d)-[:`소속`]->(co)"),
    ]

    total = 0
    details: dict[str, int] = {}
    for desc, count_q, fix_q in steps:
        try:
            rows = await run_cypher(count_q)
            cnt = int(rows[0]["cnt"]) if rows else 0
            if cnt:
                await run_cypher(fix_q)
                details[desc] = cnt
                total += cnt
                logger.info(f"[RelNorm] {desc}: {cnt}건")
        except Exception as e:
            logger.warning(f"[RelNorm] {desc} 실패 (무시): {e}")

    return {"total": total, "details": details}


@router.post("/knowledge/analyze-relationships")
async def analyze_relationships_stream(
    background_tasks: BackgroundTasks,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    background_tasks.add_task(
        _log_activity, 0, "워크메이트[supervisor→knowledge]",
        "관계도 분석·재구성", f"요청자: {current_user.name}"
    )

    async def stream():
        _collector = TokenUsageCollector()
        _tok_ctx_token = _token_collector_var.set(_collector)
        _log_id = _create_log(
            context_type="archive_analyze_stream",
            meeting_id=None,
            session_id=None,
            user_id=current_user.id,
            input_data=None,
        )
        _stream_error = None
        try:
            import asyncio as _asyncio
            # 그래프 분석과 병렬로 LLM이 작업 계획을 서술
            analysis_task = _asyncio.create_task(_analyze_graph())
            _pre_sys = (
                "업무 지식 그래프를 점검하는 AI입니다. 지금 막 전체 그래프를 분석하려 합니다. "
                "어떤 항목들을 살펴볼지 한국어로 2~3단계 간결하게 나열하세요. 번호·기호 없이."
            )
            _pre_hmn = "회의체·세션·안건·문서·구성원의 관계 전체를 점검합니다."
            async for _step in _stream_plan(_pre_sys, _pre_hmn):
                yield f"data: [PLANNING] {_step}\n\n"
            analysis = await analysis_task

            sem    = analysis["semantic"]
            struct = analysis["structural"]
            member = analysis["membership"]
            counts = analysis["counts"]

            chains    = struct.get("session_chains", [])
            lifecycle = struct.get("lifecycle_gaps", [])
            stale     = struct.get("stale_links", [])
            ag_links  = sem["agenda_links"]
            doc_links = sem["doc_links"]
            ownerless = struct["ownerless_agendas"]
            orphans   = struct["orphan_documents"]
            embeddable_orphans = [d for d in orphans if d.get("emb")]
            missing_seq = sum(len(c["missing"]) for c in chains)
            lifecycle_n = len(lifecycle)
            stale_n     = len(stale)

            # 세션→안건 미연결 건수 (원칙: 회의는 아젠다와 연결되어야 한다)
            try:
                _sa_rows = await run_cypher(
                    "MATCH (s:Session)-[:`소속`|`개최`]->(mg:Meetings) "
                    "WHERE (mg)<-[:`관할`]-(:Agenda) "
                    "  AND NOT (s)-[:`진행`|`다룸멌`|`도출`]->(:Agenda) "
                    "  AND NOT (:Agenda)-[:`발제세션`]->(s) "
                    "RETURN count(s) AS cnt"
                )
                session_no_agenda_n = _sa_rows[0].get("cnt", 0) if _sa_rows else 0
            except Exception:
                session_no_agenda_n = 0

            # ── LLM이 분석 결과를 보고 현황·계획을 스스로 서술 ──────────
            findings_text = (
                f"회의체 {counts['meetings']}개 / 세션 {counts['sessions']}개 / "
                f"안건 {counts['agendas']}개 / 문서 {counts['documents']}개 스캔 완료.\n"
                f"끊긴 세션 흐름: {missing_seq}건, "
                f"세션→안건 미연결: {session_no_agenda_n}건, "
                f"회의록→안건 미연결: {lifecycle_n}건, "
                f"불필요 연결(이월·저유사도): {stale_n}건.\n"
                f"회의 간 유사 안건: {len(ag_links)}쌍"
                + (f" (상위: [{ag_links[0]['a_title']}] ↔ [{ag_links[0]['b_title']}], "
                   f"{ag_links[0]['score']*100:.0f}%)" if ag_links else "")
                + f"\n미연결 문서: {len(doc_links)}건, "
                f"담당자 없는 안건: {len(ownerless)}건, "
                f"고아 문서: {len(orphans)}건, "
                f"소속 이슈: {len(member.get('issues', []))}건."
            )
            if chains:
                c0 = chains[0]
                findings_text += f"\n예시: '{c0['mg']}' 회의체에서 {c0['count']}개 회차 미연결."

            system_findings = (
                "당신은 업무 지식 그래프를 관리하는 AI입니다. "
                "아래 그래프 분석 수치를 보고, 현재 상태와 앞으로 할 작업을 "
                "자연스러운 한국어로 간결하게 나열하세요. "
                "각 항목은 한 줄, 총 3~6개, 마크다운·번호·기호 없이 plain text만."
            )
            async for step in _stream_plan(system_findings, findings_text):
                yield f"data: [PLANNING] {step}\n\n"

            # ── 실제 그래프 재구성 수행 ──────────────────────────────────
            actionable = (missing_seq + session_no_agenda_n + lifecycle_n + stale_n +
                          len(ag_links) + len(doc_links) +
                          len(embeddable_orphans) + len(member.get("issues", [])))

            if actionable == 0:
                result = {"actions": [], "stats": {
                    "session_links": 0, "lifecycle_links": 0, "carry_links": 0,
                    "related_agendas": 0, "doc_refs": 0, "doc_attached": 0,
                    "membership_fixed": 0, "pruned_links": 0,
                }, "advisories": {
                    "ownerless_agendas": ownerless,
                    "minuteless_sessions": struct["minuteless_sessions"],
                    "isolated_persons": struct["isolated_persons"],
                }}
            else:
                result = await knowledge_agent.reconcile_graph(analysis)
                stats = result["stats"]

                # LLM이 수행 결과를 보고 planning 단계를 서술
                done_items = [
                    label.format(stats[k])
                    for k, label in [
                        ("session_links", "끊긴 회의 흐름 {}건 → 시간순 연결"),
                        ("session_agenda_links", "세션→안건 연결 {}건 생성"),
                        ("lifecycle_links", "회의록 {}개를 관련 안건과 연결"),
                        ("carry_links", "미해결 안건 {}건 다음 회차로 이월"),
                        ("related_agendas", "회의 간 유사 안건 링크 {}건 생성"),
                        ("doc_refs", "문서 참조 링크 {}건 생성"),
                        ("doc_attached", "고아 문서 {}건 → 적합 안건 자동 연결"),
                        ("membership_fixed", "소속 무결성 {}건 보정"),
                        ("pruned_links", "불필요 연결 {}건 정리"),
                    ]
                    if stats.get(k)
                ]
                if done_items:
                    actions_text = "수행한 작업 목록:\n" + "\n".join(f"- {d}" for d in done_items)
                    system_actions = (
                        "당신은 업무 지식 그래프 관리 AI입니다. "
                        "방금 그래프에 적용한 작업들을 바탕으로, 어떤 개선이 이루어졌는지 "
                        "자연스러운 한국어로 2~4줄로 서술하세요. "
                        "마크다운·번호·기호 없이 plain text만."
                    )
                    async for step in _stream_plan(system_actions, actions_text):
                        yield f"data: [PLANNING] {step}\n\n"

            # ── 관계 방향·명칭 정규화 (항상 실행) ────────────────────
            try:
                norm = await _normalize_rel_directions()
                if norm["total"]:
                    yield f"data: [PLANNING] 관계 방향·명칭 정규화 {norm['total']}건 완료\n\n"
            except Exception:
                pass

            _hl = set()
            for l in ag_links:
                _hl.add(l.get("a_title")); _hl.add(l.get("b_title"))
            for l in doc_links:
                _hl.add(l.get("ag_title"))
            for g in lifecycle:
                _hl.add(g.get("session_title"))
            for a in result.get("actions", []):
                if a.get("highlight"):
                    _hl.add(a["highlight"])
            _hl.discard(None); _hl.discard("")
            if _hl:
                yield f"data: [HIGHLIGHT] {json.dumps(list(_hl), ensure_ascii=False)}\n\n"

            report = {**result, "counts": counts,
                      "findings": {"session_missing": missing_seq, "session_groups": len(chains),
                                   "lifecycle_gaps": lifecycle_n, "stale_links": stale_n,
                                   "agenda_links": len(ag_links), "doc_links": len(doc_links),
                                   "ownerless": len(ownerless), "orphans": len(orphans),
                                   "examples": ag_links[:5], "doc_examples": doc_links[:5],
                                   "chain_examples": chains[:3]}}
            async for chunk in knowledge_agent.summarize_relationship_analysis(report):
                yield f"data: {chunk.replace(chr(10), chr(92)+chr(110))}\n\n"

        except Exception as e:
            _stream_error = e
            yield f"data: 관계도 분석 중 오류가 발생했습니다: {str(e)}\n\n"
        except BaseException as _e:
            _stream_error = _e
            raise
        finally:
            _token_collector_var.reset(_tok_ctx_token)
            _finalize(_log_id, _collector, _stream_error, None)
        yield "data: [DONE]\n\n"

    return StreamingResponse(stream(), media_type="text/event-stream")


# ─── Knowledge Base 저장 ──────────────────────────────────────────────────────
@router.post("/knowledge/store-minutes")
async def knowledge_store_minutes(
    data: _StoreMinutesReq,
    _: models.User = Depends(get_current_user),
):
    try:
        return await knowledge_agent.store_minutes(
            title=data.title, content=data.content,
            meeting_id=data.meeting_id, session_id=data.session_id,
        )
    except Exception as e:
        return {"status": "error", "detail": str(e)}


@router.post("/knowledge/store-task")
async def knowledge_store_task(
    data: _StoreTaskReq,
    _: models.User = Depends(get_current_user),
):
    try:
        return await knowledge_agent.store_task(
            content=data.content, department=data.department,
            due_date=data.due_date, meeting_id=data.meeting_id,
        )
    except Exception as e:
        return {"status": "error", "detail": str(e)}


@router.post("/knowledge/store-report")
async def knowledge_store_report(
    data: _StoreReportReq,
    _: models.User = Depends(get_current_user),
):
    try:
        return await knowledge_agent.store_report(
            title=data.title, content=data.content,
            meeting_id=data.meeting_id, score=data.score,
        )
    except Exception as e:
        return {"status": "error", "detail": str(e)}


@router.post("/knowledge/propose-relationships", summary="Knowledge Propose Relationships")
async def knowledge_propose_relationships(
    data: _ProposeRelationshipsReq,
    _: models.User = Depends(get_current_user),  # 인증 가드 (본문에서 미사용)
):
    """Neo4j 노드 간 연결 관계를 LLM이 분석해 제안. proposal_id를 반환."""
    try:
        result = await knowledge_agent.propose_relationships(
            meeting_id=data.meeting_id,
            node_types=data.node_types,
        )
        return result
    except Exception as e:
        return {"status": "error", "detail": str(e)}


@router.post("/knowledge/confirm-relationships", summary="Knowledge Confirm Relationships")
async def knowledge_confirm_relationships(
    data: _ConfirmRelationshipsReq,
    _: models.User = Depends(get_current_user),  # 인증 가드 (본문에서 미사용)
):
    """제안된 관계를 승인(Neo4j MERGE) 또는 반려(HumanJudgment 노드 생성)."""
    try:
        result = await knowledge_agent.confirm_relationships(
            proposal_id=data.proposal_id,
            approved=data.approved,
            reject_reason=data.reject_reason,
        )
        return result
    except Exception as e:
        return {"status": "error", "detail": str(e)}


# ─── 아카이브 과제 추출 ───────────────────────────────────────────────────────
@router.post("/archive/extract-agendas")
async def archive_extract_agendas(
    meeting_id: int = Form(...),
    selected_file_ids: str = Form("[]"),
    selected_similar_docs: str = Form("[]"),
    files: List[UploadFile] = File(default=[]),
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    selected_ids = json.loads(selected_file_ids) if selected_file_ids else []

    meeting = db.query(models.Meeting).filter(models.Meeting.id == meeting_id).first()
    if not meeting:
        return {"agendas": [], "error": "회의체를 찾을 수 없습니다."}

    meeting_context = _get_meeting_context(db, meeting_id)
    org_dept_pairs = _get_member_org_depts(db, meeting_id)
    previous_minutes = _get_previous_minutes(db, meeting_id)[:3]

    current_agendas = db.query(models.Agenda).filter(
        models.Agenda.meeting_id == meeting_id,
        models.Agenda.status == "ongoing",
    ).order_by(models.Agenda.created_at).all()

    pending_todos_text = ""
    if current_agendas:
        lines = []
        for a in current_agendas:
            dept = (a.department[0] if isinstance(a.department, list) and a.department else a.department) or "미지정"
            due = a.due_date.strftime("%Y-%m-%d") if a.due_date else "마감 미정"
            lines.append(f"- [{dept}] {a.title} (마감: {due})")
        pending_todos_text = "\n".join(lines)

    file_texts = []
    for fid in selected_ids:
        try:
            report = db.query(models.Report).filter(models.Report.id == int(fid)).first()
            raw = None
            if report and report.file_path:
                from r2_storage import is_r2_url as _is_r2, url_to_key as _r2_key, download_bytes as _r2_dl
                if _is_r2(report.file_path):
                    raw = _r2_dl(_r2_key(report.file_path))
                elif os.path.exists(report.file_path):
                    with open(report.file_path, "rb") as f:
                        raw = f.read()
            if report and raw:
                text = _extract_text_from_file(raw, report.file_name or "")
                if text.strip():
                    file_texts.append(f"[보고서: {report.file_name}]\n{text[:4000]}")
        except Exception as e:
            print(f"[DB 파일 추출 오류] {e}")

    current_minutes_texts = []  # 현재 회의록 (최우선 컨텍스트)
    for upload in files:
        if not upload or not upload.filename:
            continue
        try:
            raw = await upload.read()
            text = _extract_text_from_file(raw, upload.filename.lower())
            fname = upload.filename
            if text.strip():
                # 파일명에 "회의록" 포함 시 현재 회의록으로 분리
                if "회의록" in fname or "minutes" in fname.lower():
                    current_minutes_texts.append(text[:4000])
                else:
                    file_texts.append(f"[첨부: {fname}]\n{text[:4000]}")
            else:
                file_texts.append(f"[첨부: {fname}] - 텍스트 추출 불가")
        except Exception as e:
            print(f"[업로드 파일 추출 오류] {upload.filename}: {e}")

    # 현재 회의록을 이전 회의록보다 앞에 배치 (가장 최신 = 가장 높은 우선순위)
    all_minutes = current_minutes_texts + previous_minutes

    context_parts = [f"[회의체 정보]\n{meeting_context}"]
    if meeting.guidelines:
        context_parts.append(f"[회의 지침]\n{meeting.guidelines}")
    if all_minutes:
        context_parts.append(
            "[최근 회의록]\n" + "\n\n".join(f"[회의록 {i+1}]\n{m}" for i, m in enumerate(all_minutes))
        )
    if pending_todos_text:
        context_parts.append(f"[미완료 과제]\n{pending_todos_text}")
    if file_texts:
        context_parts.append("[첨부 자료]\n" + "\n\n".join(file_texts))

    org_dept_list = (
        "\n".join(
            f"- {p['company']} / {p['department']}" if p.get("company") else f"- {p['department']}"
            for p in org_dept_pairs
        ) if org_dept_pairs else "정보 없음"
    )

    try:
        parsed = await task_agent.extract_agendas_from_context(context_parts, org_dept_list, user_id=current_user.id)

        # ── draft 즉시 저장 + AgentLog ────────────────────────────────────
        import uuid as _uuid
        from datetime import datetime as _dt
        agendas_raw = parsed.get("agendas", [])
        draft_ids: list[int | None] = [None] * len(agendas_raw)
        agent_log_id: int | None = None
        try:
            for idx, ag in enumerate(agendas_raw):
                title = ag.get("title", "").strip()
                if not title:
                    continue
                due_val = None
                if ag.get("due_date"):
                    try:
                        due_val = _dt.strptime(ag["due_date"], "%Y-%m-%d")
                    except Exception:
                        pass
                dept_raw = ag.get("department")
                dept_json = [dept_raw] if dept_raw and dept_raw != "null" else None
                db_agenda = models.Agenda(
                    meeting_id=meeting_id,
                    title=title,
                    status="draft",
                    department=dept_json,
                    due_date=due_val,
                    ai_evidence=json.dumps({
                        "reasoning": ag.get("reasoning") or "",
                        "company": ag.get("company") or ag.get("organization"),
                        "start_date": ag.get("start_date"),
                    }, ensure_ascii=False),
                )
                db.add(db_agenda)
                db.flush()
                draft_ids[idx] = db_agenda.id

            agent_log = models.AgentLog(
                task_id=str(_uuid.uuid4()),
                context_type="agenda_extraction",
                meeting_id=meeting_id,
                user_id=current_user.id,
                status="success",
                input_data={
                    "selected_file_ids": selected_ids,
                    "file_count": len(file_texts),
                    "prev_minutes_count": len(previous_minutes),
                },
                output_data={
                    "agenda_count": sum(1 for d in draft_ids if d),
                    "agenda_titles": [ag.get("title") for ag in agendas_raw],
                },
                ended_at=_dt.utcnow(),
            )
            db.add(agent_log)
            db.flush()
            agent_log_id = agent_log.id
            db.commit()
            # ── draft Agenda 를 Neo4j에 즉시 동기화 (background) ────────────────────────
            try:
                import asyncio as _asyncio
                from neo4j_sync import sync_agenda as _sync_ag
                for idx, ag_raw_id in enumerate(draft_ids):
                    if ag_raw_id:
                        _asyncio.ensure_future(_sync_ag(
                            agenda_id=ag_raw_id,
                            meeting_id=meeting_id,
                            title=agendas_raw[idx].get("title", ""),
                            status="draft",
                        ))
            except Exception as _se:
                logger.warning(f"[extract-agendas] Neo4j draft sync 실패: {_se}")
        except Exception as e:
            db.rollback()
            logger.warning(f"[archive/extract-agendas] draft 저장 실패: {e}")

        # ── 컨텍스트 파일 pending → approved ─────────────────────────
        try:
            for fid in selected_ids:
                report = db.query(models.Report).filter(
                    models.Report.id == int(fid),
                    models.Report.human_status == "pending",
                ).first()
                if report:
                    report.human_status = "approved"
            db.commit()
        except Exception as _e:
            logger.warning(f"[extract-agendas] 파일 상태 업데이트 실패: {_e}")

        return {
            "agent_log_id": agent_log_id,
            "agendas": [
                {
                    "title": ag.get("title", ""),
                    "company": ag.get("company") or ag.get("organization"),
                    "department": ag.get("department"),
                    "start_date": ag.get("start_date"),
                    "due_date": ag.get("due_date"),
                    "db_id": draft_ids[idx],
                    "_state": None,
                    "_editing": False,
                }
                for idx, ag in enumerate(agendas_raw)
            ],
            "context_used": {
                "minutes_count": len(previous_minutes),
                "current_agendas_count": len(current_agendas),
                "files_count": len(file_texts),
            },
        }
    except Exception as e:
        print(f"[archive/extract-agendas 오류] {e}")
        return {"agendas": [], "error": f"AI 분석 중 오류: {str(e)}"}


# ─── 아카이브 채팅 기반 과제 업데이트 ────────────────────────────────────────
@router.post("/archive/chat-extract")
async def archive_chat_extract(
    data: schemas.AgentChatRequest,
    background_tasks: BackgroundTasks,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    meeting_id = data.meeting_id
    message = data.message or ""
    current_agendas = data.chat_history[0].get("agendas", []) if data.chat_history else []

    meeting_context = _get_meeting_context(db, meeting_id) if meeting_id else ""
    org_dept_pairs = _get_member_org_depts(db, meeting_id) if meeting_id else []
    org_dept_list = (
        "\n".join(
            f"- {p['company']} / {p['department']}" if p.get("company") else f"- {p['department']}"
            for p in org_dept_pairs
        ) if org_dept_pairs else "정보 없음"
    )
    current_agendas_text = json.dumps(current_agendas, ensure_ascii=False, indent=2) if current_agendas else "없음"

    async def stream():
        _collector = TokenUsageCollector()
        _tok_ctx_token = _token_collector_var.set(_collector)
        _log_id = _create_log(
            context_type="agenda_extraction",
            meeting_id=meeting_id or None,
            session_id=None,
            user_id=current_user.id,
            input_data={"message": message[:300]},
        )
        _stream_error = None
        try:
            cnt = len(current_agendas)
            # LLM이 요청 내용을 보고 처리 계획을 스스로 서술
            _plan_sys = (
                "업무 과제 관리 AI입니다. 사용자 요청을 바탕으로 과제 목록을 어떻게 처리할지 "
                "한국어로 2~3단계를 간결하게 나열하세요. 각 단계는 짧은 한 문장, 번호·기호 없이."
            )
            _plan_hmn = f"현재 과제 {cnt}건. 사용자 요청: {message[:300]}"
            async for _step in _stream_plan(_plan_sys, _plan_hmn):
                yield f"data: [PLANNING] {_step}\n\n"

            parsed = await task_agent.chat_update_agendas(message, meeting_context, org_dept_list, current_agendas_text)
            if not parsed:
                parsed = {"agendas": current_agendas, "message": message}

            agendas = parsed.get("agendas", current_agendas)
            result = {
                "agendas": [
                    {
                        "title": ag.get("title", ""),
                        "company": ag.get("company") or ag.get("organization"),
                        "department": ag.get("department"),
                        "priority": ag.get("priority", "normal"),
                        "start_date": ag.get("start_date"),
                        "due_date": ag.get("due_date"),
                        "_state": None,
                        "_editing": False,
                    }
                    for ag in agendas
                ],
                "reply": parsed.get("message", "과제 목록을 업데이트했습니다."),
            }
            yield f"data: [RESULT] {json.dumps(result, ensure_ascii=False)}\n\n"
        except Exception as e:
            _stream_error = e
            logger.warning(f"[chat-extract] 오류: {e}")
            fallback = {"agendas": current_agendas, "reply": f"오류: {str(e)}"}
            yield f"data: [RESULT] {json.dumps(fallback, ensure_ascii=False)}\n\n"
        except BaseException as _e:
            _stream_error = _e
            raise
        finally:
            _token_collector_var.reset(_tok_ctx_token)
            _finalize(_log_id, _collector, _stream_error, None)
        yield "data: [DONE]\n\n"

    return StreamingResponse(stream(), media_type="text/event-stream")


# ─── 아카이브 파일 AI 검토 ────────────────────────────────────────────────────
@router.post("/archive/analyze-file")
async def analyze_archive_file(
    file: Optional[UploadFile] = File(None),
    file_name: str = Form(""),
    file_type: str = Form(""),
    dept_name: str = Form(""),
    graph_context: str = Form(""),
    candidate_agendas: str = Form("[]"),
    meeting_id: Optional[int] = Form(None),
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    # 후보 과제(JSON 문자열) 파싱
    try:
        candidate_list = json.loads(candidate_agendas) if candidate_agendas else []
    except Exception:
        candidate_list = []
    if not isinstance(candidate_list, list):
        candidate_list = []

    # 첨부 파일에서 실제 텍스트 추출 (PDF/DOCX/XLSX/텍스트)
    file_content = ""
    if file is not None:
        try:
            raw = await file.read()
            extracted = _extract_text_from_file(raw, (file.filename or file_name or "").lower())
            file_content = (extracted or "").strip()[:8000]
            if not file_content:
                file_content = "[파일에서 텍스트를 추출하지 못했습니다 — 이미지 기반 PDF일 수 있음]"
        except Exception as e:
            logger.warning(f"[analyze-file] 텍스트 추출 실패: {e}")
            file_content = f"[파일 추출 오류: {e}]"
    else:
        file_content = "[파일 미첨부 — 이름만 입력됨]"

    # LangGraph 기반 아카이브 파일 검토 에이전트 실행
    try:
        return await report_agent.analyze_archive_file(
            file_name=file_name,
            file_type=file_type,
            dept_name=dept_name,
            file_content=file_content,
            graph_context=graph_context,
            candidate_agendas=candidate_list,
            user_id=current_user.id if current_user else None,
            meeting_id=meeting_id,
        )
    except Exception as e:
        logger.warning(f"[analyze-file] LangGraph 검토 실패: {e}")
        return {
            "score": 70,
            "feedback": [f"AI 분석 중 오류: {str(e)}", "수동으로 검토해 주세요."],
            "matched_agendas": [],
            "agendas": [{"content": f"{file_name} 관련 안건 검토", "department": dept_name}],
            "related_depts": [],
        }


# ─── 아카이브 파일 AI 검토 (스트리밍) ─────────────────────────────────────────
@router.post("/archive/analyze-file/stream")
async def analyze_archive_file_stream_ep(
    file: Optional[UploadFile] = File(None),
    file_name: str = Form(""),
    file_type: str = Form(""),
    dept_name: str = Form(""),
    graph_context: str = Form(""),
    candidate_agendas: str = Form("[]"),
    meeting_id: Optional[int] = Form(None),
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    # 후보 과제(JSON 문자열) 파싱
    try:
        candidate_list = json.loads(candidate_agendas) if candidate_agendas else []
    except Exception:
        candidate_list = []
    if not isinstance(candidate_list, list):
        candidate_list = []

    # 첨부 파일에서 실제 텍스트 추출 (요청 컨텍스트 내에서 먼저 읽어둠)
    file_content = ""
    if file is not None:
        try:
            raw = await file.read()
            extracted = _extract_text_from_file(raw, (file.filename or file_name or "").lower())
            file_content = (extracted or "").strip()[:8000]
            if not file_content:
                file_content = "[파일에서 텍스트를 추출하지 못했습니다 — 이미지 기반 PDF일 수 있음]"
        except Exception as e:
            logger.warning(f"[analyze-file/stream] 텍스트 추출 실패: {e}")
            file_content = f"[파일 추출 오류: {e}]"
    else:
        file_content = "[파일 미첨부 — 이름만 입력됨]"

    async def stream():
        try:
            async for event in report_agent.analyze_archive_file_stream(
                file_name=file_name,
                file_type=file_type,
                dept_name=dept_name,
                file_content=file_content,
                graph_context=graph_context,
                candidate_agendas=candidate_list,
                user_id=current_user.id if current_user else None,
                meeting_id=meeting_id,
            ):
                yield f"data: {json.dumps(event, ensure_ascii=False)}\n\n"
        except Exception as e:
            logger.warning(f"[analyze-file/stream] 검토 실패: {e}")
            err = {
                "type": "result",
                "data": {
                    "score": 70,
                    "feedback": [f"AI 분석 중 오류: {str(e)}", "수동으로 검토해 주세요."],
                    "matched_agendas": [],
                    "agendas": [{"content": f"{file_name} 관련 안건 검토", "department": dept_name}],
                    "related_depts": [],
                },
            }
            yield f"data: {json.dumps(err, ensure_ascii=False)}\n\n"
        yield "data: [DONE]\n\n"

    return StreamingResponse(stream(), media_type="text/event-stream")


# ─── 아젠다 commit (승인→ongoing 업데이트 / 반려→삭제) ────────────────────────
@router.post("/archive/agendas/commit")
async def commit_draft_agendas(
    data: dict,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    import uuid as _uuid
    from datetime import datetime as _dt

    meeting_id: int = data.get("meeting_id", 0)
    approved: list = data.get("approved", [])   # [{db_id, assignee_name, dept, due_date}]
    rejected_ids: list = data.get("rejected_ids", [])  # [int]

    # 반려된 draft 삭제
    if rejected_ids:
        db.query(models.Agenda).filter(
            models.Agenda.id.in_(rejected_ids),
            models.Agenda.status == "draft",
        ).delete(synchronize_session=False)

    # 승인된 항목 업데이트 또는 신규 생성
    updated_ids = []
    for item in approved:
        db_id = item.get("db_id")
        agenda = db.query(models.Agenda).filter(models.Agenda.id == db_id).first() if db_id else None

        if agenda:
            # 기존 draft 업데이트
            if item.get("title"):
                agenda.title = item["title"]
            agenda.status = "ongoing"
            if item.get("dept"):
                agenda.department = [item["dept"]]
            if item.get("due_date"):
                try:
                    agenda.due_date = _dt.strptime(item["due_date"], "%Y-%m-%d")
                except Exception:
                    pass
        else:
            # db_id 없는 신규 항목 직접 생성
            if not item.get("title") or not meeting_id:
                continue
            agenda = models.Agenda(
                meeting_id=meeting_id,
                title=item["title"],
                status="ongoing",
                department=[item["dept"]] if item.get("dept") else [],
            )
            if item.get("due_date"):
                try:
                    agenda.due_date = _dt.strptime(item["due_date"], "%Y-%m-%d")
                except Exception:
                    pass
            db.add(agenda)
            db.flush()

        updated_ids.append(agenda.id)

    # AgentLog 기록
    db.add(models.AgentLog(
        task_id=str(_uuid.uuid4()),
        context_type="agenda_commit",
        meeting_id=meeting_id or None,
        user_id=current_user.id,
        status="success",
        input_data={"approved_count": len(approved), "rejected_count": len(rejected_ids)},
        output_data={"updated_ids": updated_ids, "deleted_ids": rejected_ids},
        ended_at=_dt.utcnow(),
    ))
    db.commit()

    # Neo4j 동기화: 승인된 건 그래프에 추가, 반려된 건 그래프에서 삭제
    from neo4j_sync import sync_agenda as _sync_ag, delete_agenda as _del_ag
    import json as _json
    for ag_id in updated_ids:
        ag = db.query(models.Agenda).filter(models.Agenda.id == ag_id).first()
        if ag:
            dept_str = _json.dumps(ag.department, ensure_ascii=False) if isinstance(ag.department, (dict, list)) else (ag.department or "")
            try:
                await _sync_ag(
                    ag.id, ag.meeting_id,
                    title=ag.title, status=ag.status,
                    assignee_id=ag.assignee_id,
                    priority=ag.priority or "medium",
                    due_date=ag.due_date.isoformat() + 'Z' if ag.due_date else None,
                    department=dept_str,
                )
            except Exception as e:
                logger.warning(f"[commit] Neo4j sync 실패 (agenda {ag_id}): {e}")
    for ag_id in rejected_ids:
        try:
            await _del_ag(ag_id)
        except Exception as e:
            logger.warning(f"[commit] Neo4j 삭제 실패 (agenda {ag_id}): {e}")

    return {"updated": updated_ids, "deleted": rejected_ids}


# ─── 회의체 draft 아젠다 조회 (과제추출 탭 복원용) ──────────────────────────
@router.get("/meetings/{meeting_id}/draft-agendas")
async def get_draft_agendas(
    meeting_id: int,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    agendas = (
        db.query(models.Agenda)
        .filter(models.Agenda.meeting_id == meeting_id, models.Agenda.status == "draft")
        .order_by(models.Agenda.created_at.asc())
        .all()
    )
    def _parse_ev(ev):
        if not ev: return {}
        try: return json.loads(ev)
        except: return {}

    return [
        {
            "db_id": a.id,
            "title": a.title,
            "department": a.department,
            "due_date": a.due_date.strftime("%Y-%m-%d") if a.due_date else None,
            "start_date": _parse_ev(a.ai_evidence).get("start_date"),
            "company": _parse_ev(a.ai_evidence).get("company") or _parse_ev(a.ai_evidence).get("organization"),
        }
        for a in agendas
    ]


# ─── 회의체 아젠다 목록 조회 ──────────────────────────────────────────────────
@router.get("/meetings/{meeting_id}/agendas")
async def get_meeting_agendas(
    meeting_id: int,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    agendas = (
        db.query(models.Agenda)
        .filter(models.Agenda.meeting_id == meeting_id, models.Agenda.status != "draft")
        .order_by(models.Agenda.created_at.desc())
        .all()
    )
    def _dept_str(d):
        if not d: return None
        return d[0] if isinstance(d, list) else str(d)

    return [
        {
            "id": a.id,
            "title": a.title,
            "content": a.title,
            "status": a.status,
            "department": a.department,
            "dept": _dept_str(a.department),
            "assignee_id": a.assignee_id,
            "due_date": a.due_date.strftime("%Y-%m-%d") if a.due_date else None,
            "ai_evidence": a.ai_evidence,
            "created_at": a.created_at.isoformat() + 'Z' if a.created_at else None,
        }
        for a in agendas
    ]


# ─── 아젠다 상세 수정 (제목/부서/마감일/우선순위) ───────────────────────────
@router.patch("/archive/agendas/{agenda_id}")
async def update_agenda(
    agenda_id: int,
    data: dict,
    background_tasks: BackgroundTasks,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    agenda = db.query(models.Agenda).filter(models.Agenda.id == agenda_id).first()
    if not agenda:
        from fastapi import HTTPException
        raise HTTPException(status_code=404, detail="Agenda not found")
    if "title" in data and data["title"] is not None:
        agenda.title = data["title"]
    if "department" in data:
        raw_dept = data["department"]
        agenda.department = [raw_dept] if raw_dept else None
    if "due_date" in data:
        if data["due_date"]:
            from datetime import datetime as _dt
            try:
                agenda.due_date = _dt.strptime(data["due_date"][:10], "%Y-%m-%d")
            except Exception:
                pass
        else:
            agenda.due_date = None
    if "priority" in data and data["priority"] is not None:
        agenda.priority = data["priority"]
    if "status" in data and data["status"] is not None:
        agenda.status = data["status"]
    db.commit()
    db.refresh(agenda)
    # Neo4j 동기화
    from neo4j_sync import sync_agenda as _sync_ag
    dept_str = (agenda.department[0] if isinstance(agenda.department, list) and agenda.department else (agenda.department or ""))
    background_tasks.add_task(
        _sync_ag,
        agenda_id=agenda.id,
        meeting_id=agenda.meeting_id,
        title=agenda.title,
        status=agenda.status or "ongoing",
        priority=agenda.priority or "medium",
        due_date=agenda.due_date.isoformat() if agenda.due_date else None,
        department=dept_str,
    )
    return {
        "ok": True,
        "id": agenda.id,
        "title": agenda.title,
        "department": agenda.department,
        "due_date": agenda.due_date.isoformat() if agenda.due_date else None,
        "priority": agenda.priority,
        "status": agenda.status,
    }


# ─── 아젠다 상태 변경 (완료/진행 등) ─────────────────────────────────────────
@router.patch("/archive/agendas/{agenda_id}/status")
async def update_agenda_status(
    agenda_id: int,
    data: dict,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    agenda = db.query(models.Agenda).filter(models.Agenda.id == agenda_id).first()
    if not agenda:
        from fastapi import HTTPException
        raise HTTPException(status_code=404, detail="Agenda not found")
    new_status = data.get("status", "done")
    agenda.status = new_status
    db.commit()
    return {"ok": True, "status": new_status}


# ─── 보고자료 편집 ────────────────────────────────────────────────────────────
@router.patch("/archive/reports/{report_id}")
async def update_report(
    report_id: int,
    data: dict,
    background_tasks: BackgroundTasks,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    from fastapi import HTTPException
    report = db.query(models.Report).filter(models.Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")
    if "file_name" in data and data["file_name"] is not None:
        report.file_name = data["file_name"]
    if "submitter_department" in data and data["submitter_department"] is not None:
        report.submitter_department = data["submitter_department"]
    if "human_status" in data and data["human_status"] is not None:
        report.human_status = data["human_status"]
    db.commit()
    db.refresh(report)
    from neo4j_sync import sync_report as _sync_rp
    background_tasks.add_task(
        _sync_rp,
        report_id=report.id,
        meeting_id=report.meeting_id,
        file_name=report.file_name,
        file_path=report.file_path,
        submitter_department=report.submitter_department,
        human_status=report.human_status,
        related_agenda_ids=report.related_agenda_ids or [],
    )
    return {
        "ok": True,
        "id": report.id,
        "file_name": report.file_name,
        "submitter_department": report.submitter_department,
        "human_status": report.human_status,
    }


# ─── 회의록 편집 (session_id 기반 lookup) ────────────────────────────────────
@router.patch("/archive/minutes/by-session/{session_id}")
async def update_minutes_by_session(
    session_id: int,
    data: dict,
    background_tasks: BackgroundTasks,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    from fastapi import HTTPException
    minutes = db.query(models.Minutes).filter(models.Minutes.session_id == session_id).first()
    if not minutes:
        raise HTTPException(status_code=404, detail="Minutes not found for session")
    if "file_name" in data and data["file_name"] is not None:
        minutes.file_name = data["file_name"]
    if "status" in data and data["status"] is not None:
        minutes.status = data["status"]
    db.commit()
    db.refresh(minutes)
    from neo4j_sync import sync_minutes as _sync_mn
    background_tasks.add_task(
        _sync_mn,
        minutes_id=minutes.id,
        session_id=minutes.session_id,
        file_name=minutes.file_name,
        file_path=minutes.file_path,
        recorder_id=minutes.recorder_id,
        content_summary=minutes.content_summary,
        content_original=minutes.content_original,
        status=minutes.status,
    )
    return {
        "ok": True,
        "id": minutes.id,
        "file_name": minutes.file_name,
        "status": minutes.status,
    }


# ─── 회의록 편집 ──────────────────────────────────────────────────────────────
@router.patch("/archive/minutes/{minutes_id}")
async def update_minutes(
    minutes_id: int,
    data: dict,
    background_tasks: BackgroundTasks,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    from fastapi import HTTPException
    minutes = db.query(models.Minutes).filter(models.Minutes.id == minutes_id).first()
    if not minutes:
        raise HTTPException(status_code=404, detail="Minutes not found")
    if "file_name" in data and data["file_name"] is not None:
        minutes.file_name = data["file_name"]
    if "status" in data and data["status"] is not None:
        minutes.status = data["status"]
    db.commit()
    db.refresh(minutes)
    from neo4j_sync import sync_minutes as _sync_mn
    background_tasks.add_task(
        _sync_mn,
        minutes_id=minutes.id,
        session_id=minutes.session_id,
        file_name=minutes.file_name,
        file_path=minutes.file_path,
        recorder_id=minutes.recorder_id,
        content_summary=minutes.content_summary,
        content_original=minutes.content_original,
        status=minutes.status,
    )
    return {
        "ok": True,
        "id": minutes.id,
        "file_name": minutes.file_name,
        "status": minutes.status,
    }


# ─── 아젠다 삭제 ──────────────────────────────────────────────────────────────
@router.delete("/archive/agendas/{agenda_id}")
async def delete_agenda_item(
    agenda_id: int,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    db.query(models.Agenda).filter(models.Agenda.id == agenda_id).delete(synchronize_session=False)
    db.commit()
    from neo4j_sync import delete_agenda as _del_ag
    try:
        await _del_ag(agenda_id)
    except Exception:
        pass
    return {"ok": True}


# ─── HITL 검토 저장 ──────────────────────────────────────────────────────────
class HitlReviewCreate(BaseModel):
    target_type: str
    agenda_id: Optional[int] = None
    report_id: Optional[int] = None
    agent_log_id: Optional[int] = None
    status: str = "edited"
    comment: Optional[str] = None


@router.post("/hitl-reviews")
async def create_hitl_review(
    data: HitlReviewCreate,
    background_tasks: BackgroundTasks,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    review = models.HitlReview(
        agent_log_id=data.agent_log_id,
        target_type=data.target_type,
        agenda_id=data.agenda_id,
        report_id=data.report_id,
        status=data.status,
        reviewer_id=current_user.id,
        comment=data.comment,
        reviewed_at=datetime.utcnow(),
    )
    db.add(review)
    db.commit()
    db.refresh(review)

    # ── Neo4j HumanJudgment 즉시 동기화 ──────────────────────────────────────
    try:
        from neo4j_sync import sync_human_judgment as _sync_hj
        meeting_id_for_sync: int | None = None
        if review.target_type == "agenda" and review.agenda_id:
            _ag = db.query(models.Agenda).filter(models.Agenda.id == review.agenda_id).first()
            if _ag:
                meeting_id_for_sync = _ag.meeting_id
        background_tasks.add_task(
            _sync_hj,
            review_id=review.id,
            meeting_id=meeting_id_for_sync,
            judgment=review.status,
            reason=review.comment,
            target_type=review.target_type,
            target_id=review.agenda_id or review.report_id,
            judged_at=review.reviewed_at.isoformat() if review.reviewed_at else None,
            created_at=review.created_at.isoformat() if review.created_at else None,
        )
    except Exception as _se:
        logger.warning(f"[hitl-reviews] Neo4j HumanJudgment sync 실패: {_se}")

    return {"id": review.id, "status": review.status}


class HitlReviewPatch(BaseModel):
    status: Optional[str] = None
    comment: Optional[str] = None


@router.patch("/hitl-reviews/{hj_id}")
async def update_hitl_review(
    hj_id: int,
    data: HitlReviewPatch,
    background_tasks: BackgroundTasks,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    from fastapi import HTTPException
    review = db.query(models.HitlReview).filter(models.HitlReview.id == hj_id).first()
    if not review:
        raise HTTPException(status_code=404, detail="HitlReview not found")
    if data.status is not None:
        review.status = data.status
    if data.comment is not None:
        review.comment = data.comment
    review.reviewed_at = datetime.utcnow()
    review.reviewer_id = current_user.id
    db.commit()
    db.refresh(review)

    try:
        from neo4j_sync import sync_human_judgment as _sync_hj
        meeting_id_for_sync: int | None = None
        if review.target_type == "agenda" and review.agenda_id:
            _ag = db.query(models.Agenda).filter(models.Agenda.id == review.agenda_id).first()
            if _ag:
                meeting_id_for_sync = _ag.meeting_id
        background_tasks.add_task(
            _sync_hj,
            review_id=review.id,
            meeting_id=meeting_id_for_sync,
            judgment=review.status,
            reason=review.comment,
            target_type=review.target_type,
            target_id=review.agenda_id or review.report_id,
            judged_at=review.reviewed_at.isoformat() if review.reviewed_at else None,
            created_at=review.created_at.isoformat() if review.created_at else None,
        )
    except Exception as _se:
        logger.warning(f"[hitl-reviews] Neo4j HumanJudgment sync 실패: {_se}")

    return {"id": review.id, "status": review.status, "reviewed_at": review.reviewed_at.isoformat() if review.reviewed_at else None}


# ─── 보고서 종합 검토 (스트리밍) ──────────────────────────────────────────────
class GlobalReviewRequest(BaseModel):
    meeting_id: int
    reports_info: List[dict]
    chat_history: Optional[List[dict]] = []


@router.post("/reports/global-review/stream")
async def global_review_stream_ep(
    data: GlobalReviewRequest,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    meeting_context = _get_meeting_context(db, data.meeting_id)

    async def stream():
        try:
            async for chunk in report_agent.global_review_stream(
                reports_info=data.reports_info,
                chat_history=data.chat_history or [],
                meeting_id=data.meeting_id,
                meeting_context=meeting_context,
            ):
                yield f"data: {chunk.replace(chr(10), chr(92)+chr(110))}\n\n"
        except Exception as e:
            yield f"data: [오류] {str(e)}\n\n"
        yield "data: [DONE]\n\n"

    return StreamingResponse(stream(), media_type="text/event-stream")


# ─── 보고서 단건 검토 ─────────────────────────────────────────────────────────
class ReviewReportRequest(BaseModel):
    report_content: str
    agenda: Optional[str] = ""


@router.post("/reports/review")
async def review_report_ep(
    data: ReviewReportRequest,
    current_user: models.User = Depends(get_current_user),
):
    result = await report_agent.review_report(
        report_content=data.report_content,
        agenda=data.agenda or "",
    )
    return result


# ─── 보고서 HITL 검토 시작 ────────────────────────────────────────────────────
class StartReportReviewRequest(BaseModel):
    thread_id: str
    report_content: str
    agenda: Optional[str] = ""


@router.post("/reports/review/start")
async def start_report_review_ep(
    data: StartReportReviewRequest,
    current_user: models.User = Depends(get_current_user),
):
    result = await report_agent.start_report_review(
        thread_id=data.thread_id,
        report_content=data.report_content,
        agenda=data.agenda or "",
    )
    return result


# ─── 보고서 HITL 검토 확정 ────────────────────────────────────────────────────
class ConfirmReportReviewRequest(BaseModel):
    thread_id: str
    approved: bool
    title: Optional[str] = ""
    content: Optional[str] = ""
    meeting_id: Optional[int] = None


@router.post("/reports/review/confirm")
async def confirm_report_review_ep(
    data: ConfirmReportReviewRequest,
    current_user: models.User = Depends(get_current_user),
):
    result = await report_agent.confirm_report_review(
        thread_id=data.thread_id,
        approved=data.approved,
        title=data.title or "",
        content=data.content or "",
        meeting_id=data.meeting_id,
    )
    return result


# ─── 과제 추출 HITL 시작 ──────────────────────────────────────────────────────
class StartExtractionRequest(BaseModel):
    thread_id: str
    content: str
    org_dept_list: Optional[str] = ""


@router.post("/extraction/review/start")
async def start_extraction_review_ep(
    data: StartExtractionRequest,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    result = await task_agent.start_extraction_review(
        thread_id=data.thread_id,
        content=data.content,
        org_dept_list=data.org_dept_list or "",
    )
    return result


# ─── 과제 추출 HITL 확정 ──────────────────────────────────────────────────────
class ConfirmExtractionRequest(BaseModel):
    thread_id: str
    approved: bool
    meeting_id: Optional[int] = None


@router.post("/extraction/review/confirm")
async def confirm_extraction_review_ep(
    data: ConfirmExtractionRequest,
    current_user: models.User = Depends(get_current_user),
):
    result = await task_agent.confirm_extraction_review(
        thread_id=data.thread_id,
        approved=data.approved,
        meeting_id=data.meeting_id,
    )
    return result


# ─── 지식 관리 채팅 (스트리밍) ───────────────────────────────────────────────
class KnowledgeChatRequest(BaseModel):
    message: str
    chat_history: Optional[List[dict]] = []
    meeting_id: Optional[int] = 0


@router.post("/knowledge/chat/stream")
async def knowledge_chat_stream_ep(
    data: KnowledgeChatRequest,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    meeting_context = _get_meeting_context(db, data.meeting_id) if data.meeting_id else ""

    async def stream():
        try:
            async for chunk in knowledge_agent.chat_stream(
                message=data.message,
                chat_history=data.chat_history or [],
                meeting_id=data.meeting_id or 0,
                meeting_context=meeting_context,
            ):
                yield f"data: {chunk.replace(chr(10), chr(92)+chr(110))}\n\n"
        except Exception as e:
            yield f"data: [오류] {str(e)}\n\n"
        yield "data: [DONE]\n\n"

    return StreamingResponse(stream(), media_type="text/event-stream")

