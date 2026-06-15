"""Supervisor 공용 서비스 헬퍼 (P3A-4 — supervisor.py에서 분리).

DB 컨텍스트 빌더·활동 로그·파일 텍스트 추출·플랜 스트리밍 등
여러 라우터가 공유하는 비-HTTP 로직. (routers/는 HTTP 변환만 담당)
"""

import logging
import uuid
from datetime import datetime
from typing import List, cast

from langchain_core.messages import HumanMessage, SystemMessage
from sqlalchemy.orm import Session

from db import models
from db.database import SessionLocal
from routers.prompts import make_llm

logger = logging.getLogger(__name__)


def _log_activity(meeting_id: int, agent: str, action: str, detail: str = ""):
    if not meeting_id:
        return
    db = SessionLocal()
    try:
        log = models.AgentLog(
            task_id=str(uuid.uuid4()),
            context_type=f"agent_{agent.lower()}",
            meeting_id=meeting_id,
            status="success",
            output_data={"agent": agent, "action": action, "detail": detail},
            ended_at=datetime.utcnow(),
        )
        db.add(log)
        db.commit()
    except Exception as e:
        db.rollback()
        logger.warning(f"[ActivityLog Error] {e}")
    finally:
        db.close()


def _get_meeting_context(db: Session, meeting_id: int) -> str:
    if not meeting_id:
        return ""
    meeting = db.query(models.Meeting).filter(models.Meeting.id == meeting_id).first()
    if not meeting:
        return ""
    lines = [f"회의체명: {meeting.title}"]
    if meeting.description:
        lines.append(f"회의 목적: {meeting.description}")
    members = (
        db.query(models.MeetingMember)
        .filter(models.MeetingMember.meeting_id == meeting_id)
        .all()
    )
    if members:
        member_parts = []
        for m in members:
            user = db.query(models.User).filter(models.User.id == m.user_id).first()
            if user:
                role_label = "운영자" if m.meeting_role == "admin" else "참여자"
                dept = user.department or ""
                member_parts.append(f"{user.name}({dept}, {role_label})")
        if member_parts:
            lines.append(f"참여자: {', '.join(member_parts)}")
    return "\n".join(lines)


def _get_member_org_depts(db: Session, meeting_id: int) -> List[dict]:
    """Return unique (company, department) pairs from meeting members."""
    from sqlalchemy.orm import joinedload

    members = (
        db.query(models.MeetingMember)
        .options(joinedload(models.MeetingMember.user))
        .filter(models.MeetingMember.meeting_id == meeting_id)
        .all()
    )
    seen: set = set()
    result: List[dict] = []
    for m in members:
        if not m.user or not m.user.department:
            continue
        company = m.user.company_name or ""
        dept = m.user.department
        if (company, dept) not in seen:
            seen.add((company, dept))
            result.append({"company": company, "department": dept})
    return result


def _get_previous_minutes(db: Session, meeting_id: int) -> List[str]:
    sessions = (
        db.query(models.MeetingSession)
        .filter(
            models.MeetingSession.meeting_id == meeting_id,
            models.MeetingSession.status.in_(["ended", "ENDED"]),
        )
        .order_by(models.MeetingSession.ended_at.desc())
        .all()
    )

    result = []
    for s in sessions:
        if not s.minutes:
            continue
        if s.minutes.content_summary:
            result.append(s.minutes.content_summary)
        elif s.minutes.file_path:
            # content_summary 없으면 파일에서 직접 텍스트 추출 (STT 원문 등)
            try:
                from storage.r2_storage import (
                    is_r2_url as _is_r2,
                    url_to_key as _r2_key,
                    download_bytes as _r2_dl,
                )

                raw = (
                    _r2_dl(_r2_key(s.minutes.file_path))
                    if _is_r2(s.minutes.file_path)
                    else None
                )
                if raw:
                    text = _extract_text_from_file(
                        raw, s.minutes.file_name or "minutes.pdf"
                    )
                    if text.strip():
                        result.append(text[:3000])
            except Exception:
                pass
    return result


def _format_schedule_table(table: list) -> str:
    """pdfplumber 테이블을 팀별 주차 일정 구조화 텍스트로 변환.

    pdfplumber가 셀을 세로로 분리해서 읽는 경우(팀명 행과 작업 행이 별도 행으로 추출)를
    처리하기 위해 row[0]이 None/빈값인 행을 직전 팀에 병합한다.
    """
    if not table or len(table) < 2:
        return ""

    def clean(cell) -> str:
        if cell is None:
            return ""
        return " ".join(str(cell).split())

    header = [clean(c) for c in table[0]]
    header_joined = " ".join(header)

    if not any(k in header_joined for k in ["주", "팀", "담당", "안건"]):
        return ""

    # 팀별로 col_idx → 작업명 매핑 (같은 팀의 여러 행을 병합)
    team_order: list = []
    team_tasks: dict = {}
    current_team: str | None = None

    for row in table[1:]:
        team_cell = clean(row[0]) if len(row) > 0 else ""
        if team_cell:
            current_team = team_cell
            if current_team not in team_tasks:
                team_order.append(current_team)
                team_tasks[current_team] = {}
        if current_team is None:
            continue
        for col_idx in range(1, len(row)):
            cell = clean(row[col_idx])
            if cell and col_idx not in team_tasks[current_team]:
                team_tasks[current_team][col_idx] = cell

    lines = ["[팀별 업무 일정표]"]
    for team in team_order:
        lines.append(f"[{team}]")
        tasks = team_tasks[team]
        if tasks:
            for col_idx in sorted(tasks):
                col_label = (
                    header[col_idx]
                    if col_idx < len(header) and header[col_idx]
                    else f"열{col_idx}"
                )
                lines.append(f"  {col_label}: {tasks[col_idx]}")
        else:
            lines.append("  (배정된 작업 없음)")

    return "\n".join(lines) if len(lines) > 1 else ""


def _extract_text_from_file(raw: bytes, filename: str) -> str:
    import io

    if filename.endswith(".pdf"):
        # 1순위: pdfplumber — 표는 구조화 추출, 나머지는 텍스트 추출
        try:
            import pdfplumber

            parts = []
            with pdfplumber.open(io.BytesIO(raw)) as pdf:
                for page in pdf.pages:
                    # 표 구조화 추출
                    tables = page.extract_tables()
                    for t in tables or []:
                        formatted = _format_schedule_table(t)
                        if formatted:
                            parts.append(formatted)
                    # 일반 텍스트 추출 (표 외 영역 포함)
                    page_text = page.extract_text() or ""
                    if page_text.strip():
                        parts.append(page_text.strip())
            result = "\n\n".join(parts).strip()
            if result:
                return result
        except ModuleNotFoundError:
            pass
        except Exception as e:
            logger.warning(f"[extract] pdfplumber 실패, pypdf로 대체: {e}")
        # 2순위: pypdf (기본 의존성)
        try:
            import pypdf

            reader = pypdf.PdfReader(io.BytesIO(raw))
            pages = [page.extract_text() or "" for page in reader.pages]
            return "\n".join(pages).strip()
        except Exception as e:
            return f"[PDF 추출 오류: {e}]"

    if filename.endswith(".docx"):
        try:
            import docx as _docx

            doc = _docx.Document(io.BytesIO(raw))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception as e:
            return f"[DOCX 추출 오류: {e}]"

    if filename.endswith((".xlsx", ".xls")):
        try:
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    row_text = "\t".join(str(c) if c is not None else "" for c in row)
                    if row_text.strip():
                        lines.append(row_text)
            return "\n".join(lines)
        except Exception as e:
            return f"[XLSX 추출 오류: {e}]"

    for enc in ("utf-8", "euc-kr", "cp949"):
        try:
            return raw.decode(enc)
        except Exception:
            pass
    return ""


async def _stream_plan(system_content: str, human_content: str):
    """LLM이 작업 계획을 줄 단위로 스트리밍 생성합니다."""
    llm = make_llm(temperature=0.2, streaming=True)
    buf = ""
    async for chunk in llm.astream(
        [
            SystemMessage(content=system_content),
            HumanMessage(content=human_content),
        ]
    ):
        if chunk.content:
            buf += cast(str, chunk.content)
            while "\n" in buf:
                line, buf = buf.split("\n", 1)
                line = line.strip().lstrip("-•·▪▸◦*0123456789.").strip()
                if line:
                    yield line
    tail = buf.strip().lstrip("-•·▪▸◦*0123456789.").strip()
    if tail:
        yield tail


# ─── 세션 컨텍스트 빌더 ───────────────────────────────────────────────────────


def _build_session_context(db: Session, session_id: int) -> dict:
    """session_id 기준으로 DB 데이터를 조립해 반환합니다."""
    session = (
        db.query(models.MeetingSession)
        .filter(models.MeetingSession.id == session_id)
        .first()
    )
    if not session:
        return {}

    # 참석자 (SessionMember → User 조인)
    member_rows = (
        db.query(models.SessionMember)
        .filter(models.SessionMember.session_id == session_id)
        .all()
    )
    user_ids = [m.user_id for m in member_rows]
    role_map = {m.user_id: m.role for m in member_rows}
    users = db.query(models.User).filter(models.User.id.in_(user_ids)).all() if user_ids else []

    # 안건 (SessionAgenda → Agenda)
    sa_rows = (
        db.query(models.SessionAgenda)
        .filter(models.SessionAgenda.session_id == session_id)
        .all()
    )
    agenda_ids = [sa.agenda_id for sa in sa_rows]
    agendas = (
        db.query(models.Agenda).filter(models.Agenda.id.in_(agenda_ids)).all()
        if agenda_ids
        else []
    )

    # 부모 회의체
    meeting = (
        db.query(models.Meeting)
        .filter(models.Meeting.id == session.meeting_id)
        .first()
    )

    # 회의록
    minutes = (
        db.query(models.Minutes)
        .filter(models.Minutes.session_id == session_id)
        .first()
    )

    # 실시간 요약 블록 (ongoing 상태에서 활용)
    summary_blocks = (
        db.query(models.SessionSummaryBlock)
        .filter(models.SessionSummaryBlock.session_id == session_id)
        .order_by(models.SessionSummaryBlock.block_index)
        .all()
    )

    return {
        "session": session,
        "meeting": meeting,
        "members": [(u, role_map.get(u.id, "member")) for u in users],
        "agendas": agendas,
        "minutes": minutes,
        "summary_blocks": summary_blocks,
    }


def _format_session_context_str(ctx: dict) -> str:
    """조립된 세션 컨텍스트를 LLM 프롬프트용 텍스트로 변환합니다."""
    if not ctx:
        return ""

    parts = []
    session = ctx["session"]

    # 회의체 정보
    meeting = ctx.get("meeting")
    if meeting:
        mg_lines = [f"회의체명: {meeting.title}"]
        if meeting.description:
            mg_lines.append(f"설명: {meeting.description}")
        if meeting.type:
            mg_lines.append(f"유형: {meeting.type}")
        parts.append("[회의체 정보]\n" + "\n".join(mg_lines))

    # 기본 정보
    lines = [f"회의명: {session.title}"]
    if session.scheduled_at:
        lines.append(f"일정: {session.scheduled_at.strftime('%Y-%m-%d %H:%M')}")
    if session.started_at:
        lines.append(f"시작: {session.started_at.strftime('%Y-%m-%d %H:%M')}")
    if session.ended_at:
        lines.append(f"종료: {session.ended_at.strftime('%Y-%m-%d %H:%M')}")
    if session.location:
        lines.append(f"장소: {session.location}")
    lines.append(f"상태: {session.status}")
    parts.append("[회의 기본 정보]\n" + "\n".join(lines))

    # 참석자
    if ctx["members"]:
        member_parts = [
            f"- {u.name}({u.department or ''}), 역할: {role}"
            for u, role in ctx["members"]
        ]
        parts.append("[참석자]\n" + "\n".join(member_parts))

    # 안건
    if ctx["agendas"]:
        agenda_parts = [
            f"- {a.title} (상태: {a.status})" for a in ctx["agendas"]
        ]
        parts.append("[안건]\n" + "\n".join(agenda_parts))

    # 회의록 전문 (archived일 때만 존재) — 다음 회의 아젠다 등 끝 섹션 누락 방지
    minutes = ctx.get("minutes")
    if minutes and minutes.content_summary:
        parts.append("[회의록]\n" + minutes.content_summary[:8000])

    # 실시간 요약 블록 — summary_text_override가 있으면 그걸 우선 사용 (rolling summary)
    summary_text = ctx.get("summary_text_override")
    if summary_text:
        parts.append(f"[실시간 요약 블록]\n{summary_text}")
    else:
        summary_blocks = ctx.get("summary_blocks", [])
        if summary_blocks:
            block_parts = []
            for b in summary_blocks:
                bullets = (
                    "\n".join(f"  • {bl}" for bl in (b.bullets or []))
                    if b.bullets
                    else ""
                )
                block_parts.append(f"[{b.title}]\n{bullets}" if bullets else f"[{b.title}]")
            parts.append("[실시간 요약 블록]\n" + "\n\n".join(block_parts))

    return "\n\n".join(parts)
